#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SMC_evolutions.py — Version déplacements uniquement

Objectif :
  - Parcourir les fichiers SMC et extraire uniquement les données de déplacements
  - Conserver la logique stricte : dernière colonne déterminée par A5 fusionnée ou 'Point bas droit'
  - Calculer périodique et cumulée pour chaque mesure
  - Exporter CSV en UTF-8 avec BOM
"""

from __future__ import annotations
import argparse, sys, re, unicodedata
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import pandas as pd
from openpyxl import load_workbook

# -------------------------------------------------------------
# Utilitaires
# -------------------------------------------------------------

def safe_mkdir(p: Path) -> None:
    p.mkdir(parents=True, exist_ok=True)

def is_smc_file(path: Path) -> bool:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return False
    return "smc" in path.name.lower()

def _norm(txt: str) -> str:
    if not isinstance(txt, str): return ""
    t = unicodedata.normalize("NFD", txt.upper())
    return "".join(ch for ch in t if unicodedata.category(ch) != "Mn")

def month_bounds(month_str: str) -> Tuple[pd.Timestamp, pd.Timestamp]:
    start = pd.to_datetime(month_str + "-01", format="%Y-%m-%d")
    end = start + pd.offsets.MonthEnd(0)
    return start, end

# -------------------------------------------------------------
# Lecture stricte Déplacements
# -------------------------------------------------------------

def read_deplacements_block_strict(xl_path: Path, sheet_name: str, month_str: str) -> Tuple[List[Dict], List[Dict]]:
    start, end = month_bounds(month_str)
    wb = load_workbook(xl_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return [], []
    ws = wb[sheet_name]

    # 1) dernière colonne utile
    last_col = None
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= 5 <= rng.max_row and rng.min_col <= 1 <= rng.max_col:
            last_col = rng.max_col
            break
    if not last_col:
        for cell in ws[8]:
            val = _norm(cell.value) if cell.value is not None else ""
            if "POINT BAS DROIT" in val:
                merged_rng = None
                for rng in ws.merged_cells.ranges:
                    if rng.min_row <= 8 <= rng.max_row and rng.min_col <= cell.column <= rng.max_col:
                        merged_rng = rng
                        break
                last_col = merged_rng.max_col if merged_rng else cell.column
                break
    if not last_col:
        for c in range(ws.max_column, 1, -1):
            if ws.cell(row=9, column=c).value not in (None, ""):
                last_col = c
                break
    if not last_col or last_col < 2:
        return [], []

    # 2) entêtes ligne 9
    headers = [ws.cell(row=9, column=col).value for col in range(2, last_col + 1)]
    headers = ["" if h is None else str(h).strip() for h in headers]

    # 3) données
    def to_float(v):
        if v is None: return None
        if isinstance(v, (int, float)): return float(v)
        s = str(v).replace("mm", "").replace("\u00A0", "").replace(" ", "").replace("\u2212", "-").replace(",", ".")
        try: return float(s)
        except: return None

    series_map: Dict[str, List[Tuple[pd.Timestamp, Optional[float]]]] = {
        h: [] for h in headers if h and _norm(h) != "DATE"
    }

    for r in range(10, ws.max_row + 1):
        raw_date = ws.cell(row=r, column=2).value
        if raw_date in (None, ""): continue
        try:
            d = pd.to_datetime(raw_date, errors="raise", dayfirst=True)
        except: continue
        for idx, h in enumerate(headers, start=2):
            if not h or _norm(h) == "DATE":
                continue
            v = to_float(ws.cell(row=r, column=idx).value)
            series_map[h].append((d, v))

    # 4) calculs
    rows_disp: List[Dict] = []
    month_dates: set = set()

    for label, pairs in series_map.items():
        s = pd.Series({d: v for d, v in pairs if v is not None}).sort_index()
        if s.empty:
            periodic = None
            cumulative = None
            outside = None
        else:
            cumulative = float(s.iloc[-1])
            s_before = s[s.index < start]
            v_before = float(s_before.iloc[-1]) if not s_before.empty else None
            s_in = s[(s.index >= start) & (s.index <= end)]
            v_in = float(s_in.iloc[-1]) if not s_in.empty else None
            periodic = None if (v_before is None or v_in is None) else (v_in - v_before)
            last_date_any = s.index.max()
            outside = not (start <= last_date_any <= end)
            month_dates |= set(s_in.index.to_pydatetime())

        rows_disp.append({
            "source_file": str(xl_path),
            "sheet": sheet_name,
            "month": month_str,
            "measure_label": label,
            "periodic_mm": periodic,
            "cumulative_mm": cumulative,
            "cumulative_outside_month": outside,
        })

    rows_dates: List[Dict] = []
    if month_dates:
        rows_dates.append({
            "source_file": str(xl_path),
            "sheet": sheet_name,
            "month": month_str,
            "dates": ";".join(sorted({d.strftime("%Y-%m-%d") for d in month_dates})),
        })

    return rows_dates, rows_disp

# -------------------------------------------------------------
# Traitement fichiers
# -------------------------------------------------------------

def process_excel_file(xl_path: Path, month_str: str) -> Tuple[List[Dict], List[Dict]]:
    rows_dates_all: List[Dict] = []
    rows_disp_all: List[Dict] = []
    try:
        from openpyxl import load_workbook
        xls = pd.ExcelFile(xl_path, engine="openpyxl")
    except Exception:
        return [], []

    for sheet in xls.sheet_names:
        try:
            name_norm = _norm(sheet)
            if "DEPLAC" in name_norm:
                r_dates, r_disp = read_deplacements_block_strict(xl_path, sheet, month_str)
                rows_dates_all.extend(r_dates)
                rows_disp_all.extend(r_disp)
        except Exception:
            continue
    return rows_dates_all, rows_disp_all

# -------------------------------------------------------------
# Scan et CSV
# -------------------------------------------------------------

def scan_and_process(root: Path, month_str: str) -> Tuple[pd.DataFrame, pd.DataFrame]:
    rows_dates: List[Dict] = []
    rows_disp: List[Dict] = []
    for p in root.rglob("*"):
        if p.is_file() and is_smc_file(p):
            d, z = process_excel_file(p, month_str)
            rows_dates.extend(d)
            rows_disp.extend(z)
    df_dates = pd.DataFrame(rows_dates, columns=["source_file", "sheet", "month", "dates"])
    df_disp = pd.DataFrame(rows_disp)
    return df_dates, df_disp

def save_csvs(df_dates: pd.DataFrame, df_disp: pd.DataFrame, out_dir: Path, month_str: str):
    safe_mkdir(out_dir)
    df_dates.to_csv(out_dir / f"dates_mesures_{month_str}.csv", index=False, encoding="utf-8-sig")
    df_disp.to_csv(out_dir / f"deplacements_{month_str}.csv", index=False, encoding="utf-8-sig")

# -------------------------------------------------------------
# Main
# -------------------------------------------------------------

def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    ap = argparse.ArgumentParser(description="Extraction déplacements SMC")
    ap.add_argument("--root", type=Path, help="Dossier racine")
    ap.add_argument("--month", type=str, help="Mois YYYY-MM")
    ap.add_argument("--out", type=Path, help="Dossier sortie")
    return ap.parse_args(argv)

def prompt_interactive():
    print("Mode interactif")
    root = Path(input("Dossier racine: ").strip().strip('"'))
    month = input("Mois YYYY-MM: ").strip()
    out = Path(input("Dossier sortie: ").strip().strip('"'))
    return root, month, out

def main(argv: Optional[List[str]] = None) -> int:
    if argv is None and len(sys.argv) == 1:
        root, month, out = prompt_interactive()
    else:
        args = parse_args(argv)
        root, month, out = args.root, args.month, args.out
    df_dates, df_disp = scan_and_process(root, month)
    save_csvs(df_dates, df_disp, out, month)
    print("Terminé")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
j