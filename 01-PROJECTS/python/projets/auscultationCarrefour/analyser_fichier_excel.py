#!/usr/bin/env python3
"""
Analyseur de fichier Excel - Extraction des informations de formatage et graphiques
"""

import openpyxl
from openpyxl.utils import get_column_letter

def analyser_fichier_excel(filename):
    """Analyse complète du fichier Excel existant"""

    print(f"🔍 ANALYSE COMPLÈTE DE: {filename}")
    print("="*60)

    try:
        wb = openpyxl.load_workbook(filename, data_only=False)

        # 1. Structure générale
        print(f"📋 STRUCTURE GÉNÉRALE:")
        print(f"   Nombre d'onglets: {len(wb.sheetnames)}")
        for i, sheet_name in enumerate(wb.sheetnames, 1):
            ws = wb[sheet_name]
            print(f"   {i:2d}. {sheet_name} ({ws.max_row} lignes, {ws.max_column} colonnes)")

        # 2. Analyse des graphiques existants
        print(f"\n📊 GRAPHIQUES EXISTANTS:")
        graphiques_total = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            if hasattr(ws, '_charts') and ws._charts:
                print(f"   📈 {sheet_name}: {len(ws._charts)} graphique(s)")
                graphiques_total += len(ws._charts)

                for i, chart in enumerate(ws._charts, 1):
                    print(f"      Graph {i}: Type={type(chart).__name__}")
                    if hasattr(chart, 'title') and chart.title:
                        print(f"                Titre={chart.title}")
                    if hasattr(chart, 'anchor'):
                        print(f"                Position={chart.anchor}")

        if graphiques_total == 0:
            print("   ℹ️ Aucun graphique détecté dans le fichier")

        # 3. Analyse de l'onglet Graphiques
        if "Graphiques" in wb.sheetnames:
            print(f"\n📊 ANALYSE ONGLET 'Graphiques':")
            ws_graph = wb["Graphiques"]

            # Regarder la structure des données
            print(f"   Dimensions: {ws_graph.max_row} lignes × {ws_graph.max_column} colonnes")

            # En-têtes importantes
            for row in range(1, min(20, ws_graph.max_row + 1)):
                for col in range(1, min(10, ws_graph.max_column + 1)):
                    cell = ws_graph.cell(row, col)
                    if cell.value and str(cell.value).strip():
                        # Chercher des mots-clés graphiques
                        val = str(cell.value).lower()
                        if any(keyword in val for keyword in ['graphique', 'chart', 'évolution', 'courbe', 'axe']):
                            print(f"   📈 {get_column_letter(col)}{row}: {cell.value}")

            # Formules intéressantes
            formules_graphiques = []
            for row in range(1, min(30, ws_graph.max_row + 1)):
                for col in range(1, min(20, ws_graph.max_column + 1)):
                    cell = ws_graph.cell(row, col)
                    if cell.value and str(cell.value).startswith('='):
                        formule = str(cell.value)
                        if len(formule) > 20:  # Formules complexes
                            formules_graphiques.append(f"{get_column_letter(col)}{row}: {formule[:50]}...")

            if formules_graphiques:
                print(f"\n   🔧 Formules complexes trouvées:")
                for formule in formules_graphiques[:5]:  # Première 5
                    print(f"      {formule}")

        # 4. Analyse du formatage des données
        print(f"\n🎨 ANALYSE DU FORMATAGE:")
        for sheet_name in ["OBSERVATIONS", "PROJECTION", "EVOLUTIONS XYZ", "EVOLUTIONS PM H V"]:
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                print(f"   📋 {sheet_name}:")

                # Vérifier les formats de cellules
                formats_uniques = set()
                couleurs_uniques = set()

                for row in range(9, min(15, ws.max_row + 1)):
                    for col in range(1, min(10, ws.max_column + 1)):
                        cell = ws.cell(row, col)
                        if cell.number_format:
                            formats_uniques.add(cell.number_format)
                        if hasattr(cell, 'fill') and cell.fill.start_color.index != '00000000':
                            couleurs_uniques.add(cell.fill.start_color.index)

                if formats_uniques:
                    print(f"      Formats numériques: {list(formats_uniques)[:3]}...")
                if couleurs_uniques:
                    print(f"      Couleurs de fond: {list(couleurs_uniques)[:3]}...")

        # 5. Analyse des colonnes importantes
        print(f"\n📊 COLONNES DE DONNÉES:")
        if "OBSERVATIONS" in wb.sheetnames:
            ws_obs = wb["OBSERVATIONS"]
            print(f"   OBSERVATIONS (ligne 9 - en-têtes):")
            entetes = []
            for col in range(3, min(15, ws_obs.max_column + 1)):
                entete = ws_obs.cell(9, col).value
                if entete:
                    entetes.append(f"{get_column_letter(col)}:{entete}")
            print(f"      {', '.join(entetes[:8])}...")

        wb.close()
        return True

    except Exception as e:
        print(f"❌ Erreur d'analyse: {e}")
        return False

def extraire_structure_graphiques(filename):
    """Extrait spécifiquement les informations sur les graphiques"""

    print(f"\n🎯 EXTRACTION STRUCTURE GRAPHIQUES")
    print("="*40)

    try:
        wb = openpyxl.load_workbook(filename)

        if "Graphiques" in wb.sheetnames:
            ws = wb["Graphiques"]

            # Chercher les plages de données pour graphiques
            print("📈 Plages de données détectées:")

            for row in range(1, min(50, ws.max_row + 1)):
                ligne_complete = []
                for col in range(1, min(35, ws.max_column + 1)):
                    val = ws.cell(row, col).value
                    if val is not None:
                        ligne_complete.append(f"{get_column_letter(col)}:{val}")

                if ligne_complete and len(ligne_complete) > 3:
                    # Cette ligne semble contenir des données
                    print(f"   Ligne {row}: {', '.join(ligne_complete[:5])}...")

        wb.close()

    except Exception as e:
        print(f"❌ Erreur: {e}")

if __name__ == "__main__":
    filename = input("Nom du fichier Excel à analyser: ")

    if analyser_fichier_excel(filename):
        extraire_structure_graphiques(filename)
        print(f"\n✅ Analyse terminée!")
        print(f"💡 Ces informations vont m'aider à reproduire votre style graphique")
    else:
        print(f"❌ Impossible d'analyser le fichier")
