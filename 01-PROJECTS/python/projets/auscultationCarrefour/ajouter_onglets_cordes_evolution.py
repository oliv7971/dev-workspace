#!/usr/bin/env python3
"""
Script pour ajouter les onglets CORDES et ÉVOLUTION
"""

import openpyxl
from openpyxl.utils import get_column_letter

def ajouter_onglet_cordes_ref1(filename):
    """Ajoute l'onglet CORDES REF1"""
    wb = openpyxl.load_workbook(filename)
    ws_cordes = wb.create_sheet("CORDES REF1")

    # En-têtes
    ws_cordes.cell(7, 1).value = "Date point 0"
    ws_cordes.cell(8, 1).value = "Date"
    ws_cordes.cell(8, 2).value = "Jours"
    ws_cordes.cell(9, 1).value = "Date"
    ws_cordes.cell(9, 2).value = "Jours"

    # Cibles (sans REF1 elle-même)
    targets = ["V1","V2","V3"] + [f"H{i}" for i in range(1,6)] + [f"B{i}" for i in range(1,6)] + [f"M{i}" for i in range(1,6)] + ["REF2"]

    start_col = 3
    for i, target in enumerate(targets):
        col = start_col + i
        ws_cordes.cell(9, col).value = f"CORDE_{target}_REF1"

        # Formules de calcul de distance 3D
        for row in range(10, 25):
            formule = f'=IFERROR(SQRT((INDEX(Observations.C:BJ,{row-9},MATCH("{target}_X",Observations.9:9,0))-INDEX(Observations.C:BJ,{row-9},MATCH("REF1_X",Observations.9:9,0)))^2+(INDEX(Observations.C:BJ,{row-9},MATCH("{target}_Y",Observations.9:9,0))-INDEX(Observations.C:BJ,{row-9},MATCH("REF1_Y",Observations.9:9,0)))^2+(INDEX(Observations.C:BJ,{row-9},MATCH("{target}_Z",Observations.9:9,0))-INDEX(Observations.C:BJ,{row-9},MATCH("REF1_Z",Observations.9:9,0)))^2),"")'
            ws_cordes.cell(row, col).value = formule

    wb.save(filename)
    wb.close()

def ajouter_onglet_cordes_ref2(filename):
    """Ajoute l'onglet CORDES REF2"""
    wb = openpyxl.load_workbook(filename)
    ws_cordes = wb.create_sheet("CORDES REF2")

    # En-têtes identiques
    ws_cordes.cell(7, 1).value = "Date point 0"
    ws_cordes.cell(8, 1).value = "Date"
    ws_cordes.cell(8, 2).value = "Jours"
    ws_cordes.cell(9, 1).value = "Date"
    ws_cordes.cell(9, 2).value = "Jours"

    # Cibles (sans REF2 elle-même)
    targets = ["V1","V2","V3"] + [f"H{i}" for i in range(1,6)] + [f"B{i}" for i in range(1,6)] + [f"M{i}" for i in range(1,6)] + ["REF1"]

    start_col = 3
    for i, target in enumerate(targets):
        col = start_col + i
        ws_cordes.cell(9, col).value = f"CORDE_{target}_REF2"

        # Formules vers REF2 au lieu de REF1
        for row in range(10, 25):
            formule = f'=IFERROR(SQRT((INDEX(Observations.C:BJ,{row-9},MATCH("{target}_X",Observations.9:9,0))-INDEX(Observations.C:BJ,{row-9},MATCH("REF2_X",Observations.9:9,0)))^2+(INDEX(Observations.C:BJ,{row-9},MATCH("{target}_Y",Observations.9:9,0))-INDEX(Observations.C:BJ,{row-9},MATCH("REF2_Y",Observations.9:9,0)))^2+(INDEX(Observations.C:BJ,{row-9},MATCH("{target}_Z",Observations.9:9,0))-INDEX(Observations.C:BJ,{row-9},MATCH("REF2_Z",Observations.9:9,0)))^2),"")'
            ws_cordes.cell(row, col).value = formule

    wb.save(filename)
    wb.close()

def ajouter_onglet_evolution(filename):
    """Ajoute l'onglet ÉVOLUTION PM H V"""
    wb = openpyxl.load_workbook(filename)
    ws_evol = wb.create_sheet("ÉVOLUTION PM H V")

    # En-têtes
    ws_evol.cell(7, 1).value = "Date point 0"
    ws_evol.cell(8, 1).value = "Date"
    ws_evol.cell(8, 2).value = "Jours"
    ws_evol.cell(9, 1).value = "Date"
    ws_evol.cell(9, 2).value = "Jours"

    # Toutes les cibles
    targets = ["V1","V2","V3"] + [f"H{i}" for i in range(1,6)] + [f"B{i}" for i in range(1,6)] + [f"M{i}" for i in range(1,6)] + ["REF1","REF2"]

    start_col = 3
    current_col = start_col

    for target in targets:
        # 3 colonnes par cible: ΔPM, ΔH, ΔV
        for coord in ['PM', 'H', 'V']:
            ws_evol.cell(9, current_col).value = f"Δ{target}_{coord}"

            # Formules: valeur actuelle - première valeur
            for row in range(10, 25):
                formule = f'=IFERROR(INDEX(PROJECTION.C:BJ,{row-9},MATCH("{target}_{coord}",PROJECTION.9:9,0))-INDEX(PROJECTION.C:BJ,1,MATCH("{target}_{coord}",PROJECTION.9:9,0)),"")'
                ws_evol.cell(row, current_col).value = formule

            current_col += 1

    wb.save(filename)
    wb.close()

if __name__ == "__main__":
    filename = input("Nom du fichier Excel: ")

    print("🚀 Ajout des onglets CORDES et ÉVOLUTION...")

    try:
        ajouter_onglet_cordes_ref1(filename)
        print("✅ Onglet CORDES REF1 ajouté")

        ajouter_onglet_cordes_ref2(filename)
        print("✅ Onglet CORDES REF2 ajouté")

        ajouter_onglet_evolution(filename)
        print("✅ Onglet ÉVOLUTION PM H V ajouté")

        print(f"\n🎉 TERMINÉ!")
        print(f"📊 3 nouveaux onglets ajoutés au fichier")

    except Exception as e:
        print(f"❌ Erreur: {e}")
