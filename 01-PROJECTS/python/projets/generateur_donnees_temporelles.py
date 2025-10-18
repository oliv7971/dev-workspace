#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Générateur de Données Temporelles pour Auscultation
==================================================

Ce script ajoute des données temporelles d'exemple dans les onglets Excel
pour créer des graphiques d'évolution significatifs avec plusieurs points.

Fonctionnalités:
- Génère des dates de campagnes de mesure (ex: mensuel)
- Simule l'évolution des déplacements avec tendances réalistes
- Remplit les colonnes A (dates) et données cibles
- Compatible avec les formules existantes

Auteur: Système d'automatisation Excel
Date: Octobre 2025
"""

import openpyxl
from datetime import datetime, timedelta
import random
import math
import sys
import os

class GenerateurDonneesTemporelles:
    """Génère des données temporelles pour l'auscultation"""

    def __init__(self, fichier_excel):
        self.fichier_excel = fichier_excel
        self.workbook = None

        # Configuration des données temporelles
        self.config_donnees = {
            "date_debut": datetime(2024, 1, 15),  # Première campagne
            "nb_campagnes": 12,  # 12 campagnes de mesure
            "intervalle_jours": 30,  # Une campagne par mois
            "cibles": 20,  # V1 à V20
        }

        # Paramètres de simulation des déplacements
        self.simulation = {
            "deplacement_max": 5.0,  # mm maximum
            "bruit": 0.5,  # Bruit de mesure en mm
            "tendance_lineaire": True,  # Tendance générale
            "cycles_saisonniers": True,  # Variations saisonnières
        }

    def charger_fichier(self):
        """Charge le fichier Excel"""
        try:
            print(f"📂 Chargement: {self.fichier_excel}")
            self.workbook = openpyxl.load_workbook(self.fichier_excel)
            print(f"✅ Fichier chargé - {len(self.workbook.sheetnames)} onglets")
            return True
        except Exception as e:
            print(f"❌ Erreur: {e}")
            return False

    def generer_dates_campagnes(self):
        """Génère la série de dates des campagnes de mesure"""
        dates = []
        date_courante = self.config_donnees["date_debut"]

        for i in range(self.config_donnees["nb_campagnes"]):
            dates.append(date_courante)
            date_courante += timedelta(days=self.config_donnees["intervalle_jours"])

        return dates

    def simuler_deplacement_cible(self, numero_cible, dates):
        """
        Simule l'évolution du déplacement d'une cible dans le temps

        Args:
            numero_cible: Numéro de la cible (1-20)
            dates: Liste des dates de campagne

        Returns:
            dict: Déplacements X, Y, Z pour chaque date
        """
        nb_points = len(dates)

        # Paramètres spécifiques à chaque cible
        random.seed(numero_cible * 42)  # Reproductible

        # Amplitude de déplacement selon la cible
        amplitude_base = random.uniform(0.5, self.simulation["deplacement_max"])

        # Direction préférentielle de la cible
        angle_pref = random.uniform(0, 2 * math.pi)

        deplacements = {"X": [], "Y": [], "Z": []}

        for i, date in enumerate(dates):
            t = i / (nb_points - 1)  # Normalisation 0-1

            # Tendance linéaire générale
            if self.simulation["tendance_lineaire"]:
                tendance = amplitude_base * t
            else:
                tendance = 0

            # Composante cyclique (saisonnière)
            if self.simulation["cycles_saisonniers"]:
                cycle = 0.3 * amplitude_base * math.sin(2 * math.pi * t * 2)  # 2 cycles
            else:
                cycle = 0

            # Bruit de mesure
            bruit_x = random.uniform(-self.simulation["bruit"], self.simulation["bruit"])
            bruit_y = random.uniform(-self.simulation["bruit"], self.simulation["bruit"])
            bruit_z = random.uniform(-self.simulation["bruit"], self.simulation["bruit"])

            # Calcul des déplacements
            deplacement_total = tendance + cycle

            x = deplacement_total * math.cos(angle_pref) + bruit_x
            y = deplacement_total * math.sin(angle_pref) + bruit_y
            z = deplacement_total * 0.7 + bruit_z  # Z moins marqué

            deplacements["X"].append(round(x, 3))
            deplacements["Y"].append(round(y, 3))
            deplacements["Z"].append(round(z, 3))

        return deplacements

    def remplir_onglet_evolutions_xyz(self):
        """Remplit l'onglet EVOLUTIONS XYZ avec les données temporelles"""
        if 'EVOLUTIONS XYZ' not in self.workbook.sheetnames:
            print("⚠️ Onglet EVOLUTIONS XYZ non trouvé")
            return False

        ws = self.workbook['EVOLUTIONS XYZ']
        print("\n📊 REMPLISSAGE ONGLET: EVOLUTIONS XYZ")

        # Générer les dates
        dates = self.generer_dates_campagnes()
        print(f"📅 {len(dates)} campagnes générées de {dates[0].strftime('%d/%m/%Y')} à {dates[-1].strftime('%d/%m/%Y')}")

        # Ligne de départ des données (après les en-têtes)
        ligne_debut = 10

        # Remplir les dates (colonne A)
        for i, date in enumerate(dates):
            ligne = ligne_debut + i
            ws[f'A{ligne}'] = date.strftime('%d/%m/%Y')
            ws[f'B{ligne}'] = i * self.config_donnees["intervalle_jours"]  # Jours depuis début

        # Générer et remplir les données pour chaque cible
        print("\n🎯 Génération des déplacements par cible:")

        for num_cible in range(1, min(self.config_donnees["cibles"] + 1, 21)):  # V1 à V20
            print(f"   V{num_cible}: ", end="", flush=True)

            # Simuler les déplacements
            deplacements = self.simuler_deplacement_cible(num_cible, dates)

            # Trouver les colonnes X, Y, Z pour cette cible
            # Colonne C = V1_X, D = V1_Y, E = V1_Z, F = V2_X, etc.
            col_base = 3 + (num_cible - 1) * 3  # C=3, F=6, I=9, etc.

            if col_base + 2 <= ws.max_column:
                for i, date in enumerate(dates):
                    ligne = ligne_debut + i

                    # Remplir X, Y, Z
                    ws.cell(row=ligne, column=col_base).value = deplacements["X"][i]      # X
                    ws.cell(row=ligne, column=col_base + 1).value = deplacements["Y"][i]  # Y
                    ws.cell(row=ligne, column=col_base + 2).value = deplacements["Z"][i]  # Z

                print(f"X:{deplacements['X'][-1]:.2f}, Y:{deplacements['Y'][-1]:.2f}, Z:{deplacements['Z'][-1]:.2f}")
            else:
                print("Colonne non trouvée")

        print(f"\n✅ Données générées pour {len(dates)} campagnes")
        return True

    def remplir_onglet_evolution_pm_h_v(self):
        """Remplit l'onglet ÉVOLUTION PM H V"""
        if 'ÉVOLUTION PM H V' not in self.workbook.sheetnames:
            print("⚠️ Onglet ÉVOLUTION PM H V non trouvé")
            return False

        ws = self.workbook['ÉVOLUTION PM H V']
        print("\n📊 REMPLISSAGE ONGLET: ÉVOLUTION PM H V")

        dates = self.generer_dates_campagnes()
        ligne_debut = 10

        # Remplir les dates
        for i, date in enumerate(dates):
            ligne = ligne_debut + i
            ws[f'A{ligne}'] = date.strftime('%d/%m/%Y')
            ws[f'B{ligne}'] = i * self.config_donnees["intervalle_jours"]

        # Générer données PM/H pour chaque cible
        for num_cible in range(1, min(self.config_donnees["cibles"] + 1, 21)):
            deplacements = self.simuler_deplacement_cible(num_cible, dates)

            # Convertir X,Y en PM,H (projection sur axe)
            pm_values = []
            h_values = []

            for x, y in zip(deplacements["X"], deplacements["Y"]):
                # PM = projection sur l'axe principal (combinaison X,Y)
                pm = x * 0.8 + y * 0.6  # Coefficients d'exemple
                h = y * 0.8 - x * 0.6   # H = perpendiculaire à PM

                pm_values.append(round(pm, 3))
                h_values.append(round(h, 3))

            # Trouver colonnes pour cette cible (structure peut varier)
            # On va chercher les colonnes en analysant les en-têtes
            for col in range(1, ws.max_column + 1):
                en_tete = ws.cell(row=9, column=col).value
                if en_tete and f'V{num_cible}_PM' in str(en_tete):
                    for i, date in enumerate(dates):
                        ligne = ligne_debut + i
                        ws.cell(row=ligne, column=col).value = pm_values[i]

                if en_tete and f'V{num_cible}_H' in str(en_tete):
                    for i, date in enumerate(dates):
                        ligne = ligne_debut + i
                        ws.cell(row=ligne, column=col).value = h_values[i]

            print(f"   V{num_cible}: PM:{pm_values[-1]:.2f}, H:{h_values[-1]:.2f}")

        return True

    def remplir_onglet_observations(self):
        """Remplit l'onglet OBSERVATIONS avec données de base"""
        if 'OBSERVATIONS' not in self.workbook.sheetnames:
            print("⚠️ Onglet OBSERVATIONS non trouvé")
            return False

        ws = self.workbook['OBSERVATIONS']
        print("\n📊 REMPLISSAGE ONGLET: OBSERVATIONS")

        dates = self.generer_dates_campagnes()
        ligne_debut = 10

        # Remplir dates et observations brutes
        for i, date in enumerate(dates):
            ligne = ligne_debut + i
            ws[f'A{ligne}'] = date.strftime('%d/%m/%Y')

            # Générer coordonnées brutes (différentes des évolutions)
            for num_cible in range(1, min(self.config_donnees["cibles"] + 1, 5)):  # Premières cibles
                col_base = 3 + (num_cible - 1) * 3

                if col_base + 2 <= ws.max_column:
                    # Coordonnées "absolues" simulées
                    x_abs = 1000 + random.uniform(-2, 2)
                    y_abs = 2000 + random.uniform(-2, 2)
                    z_abs = 100 + random.uniform(-1, 1)

                    ws.cell(row=ligne, column=col_base).value = round(x_abs, 4)
                    ws.cell(row=ligne, column=col_base + 1).value = round(y_abs, 4)
                    ws.cell(row=ligne, column=col_base + 2).value = round(z_abs, 4)

        print(f"✅ Observations générées pour {len(dates)} campagnes")
        return True

    def sauvegarder(self, fichier_sortie=None):
        """Sauvegarde le fichier avec les données temporelles"""
        if not fichier_sortie:
            base_name = os.path.splitext(self.fichier_excel)[0]
            fichier_sortie = f"{base_name}_AVEC_DONNEES_TEMPORELLES.xlsx"

        try:
            print(f"\n💾 Sauvegarde: {fichier_sortie}")
            self.workbook.save(fichier_sortie)
            print("✅ Fichier sauvegardé!")
            return fichier_sortie
        except Exception as e:
            print(f"❌ Erreur: {e}")
            return None

def main():
    """Fonction principale"""
    print("📅 GÉNÉRATEUR DE DONNÉES TEMPORELLES - AUSCULTATION")
    print("=" * 60)

    if len(sys.argv) > 1:
        fichier = sys.argv[1]
    else:
        fichier = input("📂 Fichier Excel: ").strip()

    if not fichier:
        fichier = "AUSCULTATION CARREFOUR_TEMPLATE_v4.xlsx"

    if not os.path.exists(fichier):
        print(f"❌ Fichier non trouvé: {fichier}")
        return

    # Créer le générateur
    generateur = GenerateurDonneesTemporelles(fichier)

    try:
        if not generateur.charger_fichier():
            return

        # Remplir les onglets avec données temporelles
        succes = True
        succes &= generateur.remplir_onglet_observations()
        succes &= generateur.remplir_onglet_evolutions_xyz()
        succes &= generateur.remplir_onglet_evolution_pm_h_v()

        if succes:
            fichier_final = generateur.sauvegarder()
            if fichier_final:
                print(f"\n🎊 DONNÉES TEMPORELLES GÉNÉRÉES AVEC SUCCÈS!")
                print(f"📈 Fichier: {fichier_final}")
                print(f"📅 12 campagnes de mesure simulées")
                print(f"🎯 Données pour 20 cibles V1-V20")
                print("\n💡 Vous pouvez maintenant régénérer les graphiques:")
                print(f"   python optimiseur_graphiques_20_cibles.py \"{fichier_final}\"")

    except Exception as e:
        print(f"❌ Erreur: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()
