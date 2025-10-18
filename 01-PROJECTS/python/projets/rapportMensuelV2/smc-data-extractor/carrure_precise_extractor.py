#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extracteur Carrure - Logique Métier PRÉCISE
Implémente la logique exacte:
- 2 tableaux distincts dans "Déplacements carrure" (A-AE et AH-BL)
- En-têtes lignes 8, 9, 10
- Déplacements cumulés = dernière ligne
- Déplacements périodiques = dernière mois - dernière avant mois (avec règle 6 mois)
"""

import openpyxl
import pandas as pd
from pathlib import Path
from datetime import datetime, date, timedelta
from typing import List, Dict, Any, Optional, Tuple
import logging

logger = logging.getLogger(__name__)

class CarrurePreciseExtractor:
    """Extracteur Carrure avec logique métier précise"""

    def __init__(self, target_month: str):
        """
        Initialise l'extracteur pour un mois cible

        Args:
            target_month: Format "YYYY-MM" (ex: "2025-09")
        """
        self.target_year, self.target_month = map(int, target_month.split('-'))

        # Fenêtre temporelle du mois
        from calendar import monthrange
        self.month_start = date(self.target_year, self.target_month, 1)
        last_day = monthrange(self.target_year, self.target_month)[1]
        self.month_end = date(self.target_year, self.target_month, last_day)

        # Limite 6 mois pour déplacements périodiques
        self.six_months_ago = self.month_start - timedelta(days=180)

        logger.info(f"Extracteur Carrure initialisé pour {self.month_start} à {self.month_end}")
        logger.info(f"Limite 6 mois: {self.six_months_ago}")

    def parse_date(self, date_value) -> Optional[date]:
        """Parse une date depuis Excel"""
        if pd.isna(date_value) or date_value is None:
            return None

        if isinstance(date_value, datetime):
            return date_value.date()

        if isinstance(date_value, str):
            try:
                return datetime.strptime(date_value.split()[0], '%Y-%m-%d').date()
            except:
                return None
        return None

    def clean_numeric_value(self, value) -> Optional[float]:
        """Nettoie une valeur numérique"""
        if pd.isna(value) or value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value) if not pd.isna(value) else None
        if isinstance(value, str):
            if value.startswith('#') or value.upper() in ['ND', 'NR', '-', '']:
                return None
            try:
                return float(value.replace(',', '.'))
            except:
                return None
        return None

    def extract_table_structure(self, worksheet, start_col: int, end_col: int) -> Dict[str, Any]:
        """
        Extrait la structure d'un tableau (en-têtes + données)

        Args:
            worksheet: Feuille Excel
            start_col: Colonne de début (1-based)
            end_col: Colonne de fin (1-based, inclusive)

        Returns:
            Dict avec structure et données du tableau
        """

        # Extraire les en-têtes depuis les lignes 8, 9, 10
        headers = {}

        for col in range(start_col, end_col + 1):
            # Ligne 8: Description longue (ex: "CARRURE GGS - Bas gauche")
            desc_long = worksheet.cell(8, col).value

            # Ligne 9: Code court (ex: "CAR-GGS-BG")
            code_court = worksheet.cell(9, col).value

            # Ligne 10: Type de mesure (ex: "DPM", "DH", "DZ", "N")
            type_mesure = worksheet.cell(10, col).value

            # Construire l'identifiant de la métrique
            if code_court and type_mesure:
                # Exclure les colonnes communes (Date, Jours, etc.)
                if type_mesure not in ['Jours écoulés', 'Date', 'Front']:
                    metric_id = f"{code_court}_{type_mesure}"
                    headers[col] = {
                        'metric_id': metric_id,
                        'code_court': code_court,
                        'type_mesure': type_mesure,
                        'description': desc_long or code_court,
                        'column': col
                    }

        # Extraire les données (à partir de ligne 11)
        data_rows = []
        date_col = start_col + 1  # Colonne B dans chaque tableau

        for row in range(11, worksheet.max_row + 1):
            # Date
            date_cell = worksheet.cell(row, date_col).value
            parsed_date = self.parse_date(date_cell)
            if not parsed_date:
                continue

            row_data = {'date': parsed_date}

            # Extraire les valeurs pour chaque métrique
            for col, header_info in headers.items():
                cell_value = worksheet.cell(row, col).value
                clean_value = self.clean_numeric_value(cell_value)

                if clean_value is not None:
                    row_data[header_info['metric_id']] = clean_value

            if len(row_data) > 1:  # Au moins date + 1 métrique
                data_rows.append(row_data)

        return {
            'headers': headers,
            'data': data_rows,
            'start_col': start_col,
            'end_col': end_col
        }

    def calculate_metric_evolution(self, data_rows: List[Dict], metric_id: str) -> Dict[str, Any]:
        """
        Calcule l'évolution d'une métrique selon la logique métier

        Args:
            data_rows: Liste des lignes de données
            metric_id: Identifiant de la métrique (ex: "CAR-GGS-BG_DPM")

        Returns:
            Dict avec déplacements cumulés et périodiques
        """

        # Filtrer les données qui ont cette métrique
        metric_data = []
        for row in data_rows:
            if metric_id in row and row[metric_id] is not None:
                metric_data.append({
                    'date': row['date'],
                    'value': row[metric_id]
                })

        if not metric_data:
            return {
                'metric_id': metric_id,
                'deplacement_cumule': None,
                'deplacement_periodique': None,
                'date_derniere_mesure': None,
                'date_derniere_avant_mois': None,
                'nb_mesures_total': 0,
                'nb_mesures_mois': 0
            }

        # Trier par date
        metric_data.sort(key=lambda x: x['date'])

        # 1. DÉPLACEMENT CUMULÉ = Dernière valeur mesurée (toutes périodes)
        last_measurement = metric_data[-1]
        deplacement_cumule = round(last_measurement['value'], 3)
        date_derniere_mesure = last_measurement['date']

        # 2. DÉPLACEMENT PÉRIODIQUE = Dernière mois - Dernière avant mois
        # Séparer les mesures par période
        mesures_mois = [m for m in metric_data if self.month_start <= m['date'] <= self.month_end]
        mesures_avant_mois = [m for m in metric_data if m['date'] < self.month_start]

        deplacement_periodique = None
        date_derniere_avant_mois = None

        if mesures_mois:  # Il y a eu des mesures ce mois
            derniere_mois = mesures_mois[-1]['value']

            if mesures_avant_mois:  # Il y a eu des mesures avant ce mois
                derniere_avant_mois = mesures_avant_mois[-1]
                date_derniere_avant_mois = derniere_avant_mois['date']

                # Vérifier la règle des 6 mois
                if date_derniere_avant_mois >= self.six_months_ago:
                    deplacement_periodique = round(derniere_mois - derniere_avant_mois['value'], 3)
                else:
                    # Dernière mesure avant le mois > 6 mois → périodique = vide
                    deplacement_periodique = None

        return {
            'metric_id': metric_id,
            'deplacement_cumule': deplacement_cumule,
            'deplacement_periodique': deplacement_periodique,
            'date_derniere_mesure': date_derniere_mesure,
            'date_derniere_avant_mois': date_derniere_avant_mois,
            'nb_mesures_total': len(metric_data),
            'nb_mesures_mois': len(mesures_mois)
        }

    def extract_deplacements_carrure(self, worksheet) -> List[Dict[str, Any]]:
        """
        Extrait les données de l'onglet "Déplacements carrure"
        Traite les 2 tableaux distincts: A-AE et AH-BL
        """

        logger.info("🏗️ Extraction Déplacements carrure (2 tableaux)")

        results = []

        # Tableau 1: Colonnes A à AE (1 à 31)
        logger.info("   📊 Tableau 1 (A-AE): CARRURE GGS")
        table1 = self.extract_table_structure(worksheet, 1, 31)

        # Calculer évolutions pour chaque métrique du tableau 1
        for col, header_info in table1['headers'].items():
            metric_id = header_info['metric_id']
            evolution = self.calculate_metric_evolution(table1['data'], metric_id)

            result_row = {
                'table': 'GGS',
                'code_point': header_info['code_court'],
                'type_mesure': header_info['type_mesure'],
                'description': header_info['description'],
                **evolution
            }
            results.append(result_row)

        # Tableau 2: Colonnes AH à BL (34 à 64)
        logger.info("   📊 Tableau 2 (AH-BL): CARRURE GVA")
        table2 = self.extract_table_structure(worksheet, 34, 64)

        # Calculer évolutions pour chaque métrique du tableau 2
        for col, header_info in table2['headers'].items():
            metric_id = header_info['metric_id']
            evolution = self.calculate_metric_evolution(table2['data'], metric_id)

            result_row = {
                'table': 'GVA',
                'code_point': header_info['code_court'],
                'type_mesure': header_info['type_mesure'],
                'description': header_info['description'],
                **evolution
            }
            results.append(result_row)

        logger.info(f"   ✅ {len(results)} métriques extraites (Déplacements carrure)")
        return results

    def extract_deplacements_cintres(self, worksheet) -> List[Dict[str, Any]]:
        """
        Extrait les données de l'onglet "Déplacements cintres-longerons"
        Traite 1 tableau unique
        """

        logger.info("🏗️ Extraction Déplacements cintres-longerons")

        results = []

        # Tableau unique: Colonnes A à AA (1 à 27 environ)
        table = self.extract_table_structure(worksheet, 1, 27)

        # Calculer évolutions pour chaque métrique
        for col, header_info in table['headers'].items():
            metric_id = header_info['metric_id']
            evolution = self.calculate_metric_evolution(table['data'], metric_id)

            result_row = {
                'table': 'CINTRES',
                'code_point': header_info['code_court'],
                'type_mesure': header_info['type_mesure'],
                'description': header_info['description'],
                **evolution
            }
            results.append(result_row)

        logger.info(f"   ✅ {len(results)} métriques extraites (Déplacements cintres)")
        return results

    def extract_carrure_file(self, file_path: Path) -> Dict[str, Any]:
        """
        Extrait toutes les données d'un fichier carrure

        Returns:
            Dict avec résultats des deux onglets
        """

        logger.info(f"🔍 Extraction fichier carrure: {file_path.name}")

        if not file_path.exists():
            logger.error(f"Fichier non trouvé: {file_path}")
            return {'deplacements_carrure': [], 'deplacements_cintres': []}

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)

            results = {
                'file_path': str(file_path),
                'deplacements_carrure': [],
                'deplacements_cintres': []
            }

            # Extraire "Déplacements carrure"
            if 'Déplacements carrure' in wb.sheetnames:
                ws_carrure = wb['Déplacements carrure']
                results['deplacements_carrure'] = self.extract_deplacements_carrure(ws_carrure)
            else:
                logger.warning("Onglet 'Déplacements carrure' non trouvé")

            # Extraire "Déplacements cintres-longerons"
            if 'Déplacements cintres-longerons' in wb.sheetnames:
                ws_cintres = wb['Déplacements cintres-longerons']
                results['deplacements_cintres'] = self.extract_deplacements_cintres(ws_cintres)
            else:
                logger.warning("Onglet 'Déplacements cintres-longerons' non trouvé")

            wb.close()

            total_metrics = len(results['deplacements_carrure']) + len(results['deplacements_cintres'])
            logger.info(f"✅ Extraction terminée: {total_metrics} métriques total")

            return results

        except Exception as e:
            logger.error(f"Erreur extraction {file_path}: {e}")
            return {'deplacements_carrure': [], 'deplacements_cintres': []}

    def save_to_csv(self, results: Dict[str, Any], output_folder: Path):
        """Sauvegarde les résultats en CSV"""

        output_folder.mkdir(exist_ok=True)

        # Combiner tous les résultats
        all_data = []

        # Ajouter données carrure
        for row in results['deplacements_carrure']:
            row['sheet_type'] = 'deplacements_carrure'
            all_data.append(row)

        # Ajouter données cintres
        for row in results['deplacements_cintres']:
            row['sheet_type'] = 'deplacements_cintres'
            all_data.append(row)

        if all_data:
            df = pd.DataFrame(all_data)

            # Réorganiser les colonnes
            cols_order = [
                'sheet_type', 'table', 'code_point', 'type_mesure', 'description',
                'deplacement_cumule', 'deplacement_periodique',
                'date_derniere_mesure', 'date_derniere_avant_mois',
                'nb_mesures_total', 'nb_mesures_mois', 'metric_id'
            ]

            # Garder seulement les colonnes qui existent
            existing_cols = [col for col in cols_order if col in df.columns]
            df = df[existing_cols]

            # Sauvegarder
            output_file = output_folder / f"carrure_extractions_{self.target_year}_{self.target_month:02d}.csv"
            df.to_csv(output_file, index=False, encoding='utf-8-sig')

            logger.info(f"💾 Résultats sauvés: {output_file}")
            logger.info(f"   📊 {len(all_data)} métriques extraites")

            return output_file
        else:
            logger.warning("Aucune donnée à sauvegarder")
            return None

def run_carrure_precise_extraction(root_folder: str, month: str, output_folder: str, log_callback=None) -> bool:
    """
    Fonction principale d'extraction carrure avec logique métier précise

    Args:
        root_folder: Dossier racine contenant les fichiers
        month: Mois cible format "YYYY-MM"
        output_folder: Dossier de sortie
        log_callback: Fonction de logging optionnelle

    Returns:
        True si succès, False sinon
    """

    def log(msg):
        if log_callback:
            log_callback(msg)
        else:
            print(msg)

    try:
        # Initialiser l'extracteur
        extractor = CarrurePreciseExtractor(month)

        # Chercher les fichiers carrure
        root_path = Path(root_folder)
        output_path = Path(output_folder)

        carrure_files = []
        for pattern in ["*carrure*.xlsx", "*carrure*.xlsm"]:
            carrure_files.extend(list(root_path.rglob(pattern)))

        if not carrure_files:
            log("⚠️  Aucun fichier carrure trouvé")
            return False

        log(f"🎯 {len(carrure_files)} fichier(s) carrure trouvé(s)")

        all_results = {'deplacements_carrure': [], 'deplacements_cintres': []}

        # Traiter chaque fichier
        for carrure_file in carrure_files:
            log(f"🔍 Traitement: {carrure_file.name}")

            file_results = extractor.extract_carrure_file(carrure_file)

            # Ajouter aux résultats globaux
            all_results['deplacements_carrure'].extend(file_results['deplacements_carrure'])
            all_results['deplacements_cintres'].extend(file_results['deplacements_cintres'])

        # Sauvegarder
        output_file = extractor.save_to_csv(all_results, output_path)

        if output_file:
            log(f"✅ Extraction carrure terminée avec succès")
            log(f"💾 Fichier généré: {output_file.name}")
            return True
        else:
            log("❌ Aucune donnée extraite")
            return False

    except Exception as e:
        log(f"❌ Erreur extraction carrure: {e}")
        return False

# Test
if __name__ == "__main__":
    test_file = Path(r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250902-rapport mensuel septembre\1-tableaux\GGS\03_100-GGS-GRD-double carrure-XYZ-tableau Jour-25-09-10.xlsm")
    test_month = "2025-09"
    test_output = Path(r"C:\temp\test_carrure_precise")

    extractor = CarrurePreciseExtractor(test_month)
    results = extractor.extract_carrure_file(test_file)
    extractor.save_to_csv(results, test_output)

    print("🎯 Test terminé - vérifiez le dossier de sortie")
