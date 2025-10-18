#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Analyseur précis des en-têtes Carrure (lignes 8, 9, 10)
Analyse les 2 tableaux distincts dans "Déplacements carrure"
"""

import openpyxl
from pathlib import Path

def analyze_carrure_headers():
    """Analyse précise des en-têtes carrure"""

    carrure_file = Path(r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250902-rapport mensuel septembre\1-tableaux\GGS\03_100-GGS-GRD-double carrure-XYZ-tableau Jour-25-09-10.xlsm")

    if not carrure_file.exists():
        print(f"❌ Fichier non trouvé: {carrure_file}")
        return

    print("🔍 ANALYSE PRÉCISE EN-TÊTES CARRURE")
    print("=" * 60)

    wb = openpyxl.load_workbook(carrure_file, data_only=True)

    # Analyser l'onglet "Déplacements carrure"
    if 'Déplacements carrure' in wb.sheetnames:
        ws = wb['Déplacements carrure']
        print(f"\n📊 ONGLET: 'Déplacements carrure'")
        print(f"Dimensions: {ws.max_row} lignes × {ws.max_column} colonnes")

        # Analyser les lignes d'en-têtes 8, 9, 10
        for header_row in [8, 9, 10]:
            print(f"\n🔍 LIGNE {header_row} (en-têtes):")

            # Tableau 1: Colonnes A à AE (1 à 31)
            print(f"   📋 TABLEAU 1 (A-AE):")
            for col in range(1, 32):  # A=1 à AE=31
                cell_value = ws.cell(header_row, col).value
                if cell_value:
                    col_letter = openpyxl.utils.get_column_letter(col)
                    print(f"      {col_letter}{header_row}: '{cell_value}'")

            # Tableau 2: Colonnes AH à BL (34 à 64)
            print(f"   📋 TABLEAU 2 (AH-BL):")
            for col in range(34, 65):  # AH=34 à BL=64
                cell_value = ws.cell(header_row, col).value
                if cell_value:
                    col_letter = openpyxl.utils.get_column_letter(col)
                    print(f"      {col_letter}{header_row}: '{cell_value}'")

        # Analyser quelques lignes de données pour comprendre la structure
        print(f"\n📈 APERÇU DONNÉES (lignes 11-15):")
        for data_row in range(11, 16):
            date_cell = ws.cell(data_row, 2).value  # Colonne B = date
            if date_cell:
                print(f"   Ligne {data_row}: Date = {date_cell}")

                # Vérifier quelques valeurs des deux tableaux
                val_table1 = ws.cell(data_row, 5).value  # Colonne E
                val_table2 = ws.cell(data_row, 37).value  # Colonne AK
                print(f"      Table1(E): {val_table1}, Table2(AK): {val_table2}")

    # Analyser l'onglet "Déplacements cintres-longerons"
    if 'Déplacements cintres-longerons' in wb.sheetnames:
        ws = wb['Déplacements cintres-longerons']
        print(f"\n📊 ONGLET: 'Déplacements cintres-longerons'")
        print(f"Dimensions: {ws.max_row} lignes × {ws.max_column} colonnes")

        # Analyser les lignes d'en-têtes 8, 9, 10
        for header_row in [8, 9, 10]:
            print(f"\n🔍 LIGNE {header_row} (en-têtes):")

            # Tableau unique - analyser toutes les colonnes avec contenu
            for col in range(1, min(30, ws.max_column + 1)):
                cell_value = ws.cell(header_row, col).value
                if cell_value:
                    col_letter = openpyxl.utils.get_column_letter(col)
                    print(f"      {col_letter}{header_row}: '{cell_value}'")

        # Aperçu données
        print(f"\n📈 APERÇU DONNÉES (lignes 11-15):")
        for data_row in range(11, 16):
            date_cell = ws.cell(data_row, 2).value
            if date_cell:
                print(f"   Ligne {data_row}: Date = {date_cell}")

    wb.close()
    print(f"\n✅ Analyse terminée")

if __name__ == "__main__":
    analyze_carrure_headers()
