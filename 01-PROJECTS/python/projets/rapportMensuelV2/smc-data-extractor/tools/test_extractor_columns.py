"""
Test simple pour voir exactement quelles colonnes lit l'extracteur
"""

import openpyxl
from pathlib import Path

def test_extractor_columns():
    """Test pour voir quelles colonnes lit l'extracteur"""
    
    excel_path = Path(r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250901-rapport mensuel aout\1-tableaux\GGS\03_104-GGS_SMC_C003_PM001_25_08-25.xlsm")
    
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['Convergences']
    
    # Reproduire la logique de l'extracteur
    def find_convergence_table_end(worksheet) -> int:
        """Copie de la fonction de l'extracteur"""
        try:
            # Chercher dans la ligne des en-têtes (ligne 8 ou 9)
            for row_num in [8, 9]:
                for col_idx in range(1, worksheet.max_column + 1):
                    cell_value = worksheet.cell(row_num, col_idx).value
                    if cell_value and isinstance(cell_value, str):
                        cell_text = cell_value.lower()
                        # Si on trouve des mots-clés de vitesse, on s'arrête à la colonne précédente
                        if any(keyword in cell_text for keyword in ['vitesse', 'évolution', 'evolution', 'velocit']):
                            print(f"🔍 Mot-clé 'vitesse' trouvé en ligne {row_num}, col {col_idx}: '{cell_value}'")
                            return col_idx - 1
            
            # Méthode alternative: chercher les cellules fusionnées
            for merged_range in worksheet.merged_cells.ranges:
                start_cell = worksheet.cell(merged_range.min_row, merged_range.min_col)
                if start_cell.value and isinstance(start_cell.value, str):
                    cell_text = start_cell.value.lower()
                    if any(keyword in cell_text for keyword in ['vitesse', 'évolution', 'evolution']):
                        print(f"🔍 Cellule fusionnée avec 'vitesse': {merged_range} - '{start_cell.value}'")
                        return merged_range.min_col - 1
            
            # Si aucune détection automatique
            last_metric_col = 1
            excluded_headers = ['Date', 'Temps (j)', 'jours écoulés', 'Distance au front']
            
            for col_idx in range(1, min(20, worksheet.max_column + 1)):
                header_cell = worksheet.cell(9, col_idx).value
                if header_cell:
                    header_text = str(header_cell).strip()
                    if (not any(excl.lower() in header_text.lower() for excl in excluded_headers) 
                        and len(header_text) <= 4
                        and header_text.upper() == header_text):
                        last_metric_col = col_idx
                        print(f"🔍 Métrique valide trouvée col {col_idx}: '{header_text}'")
            
            print(f"🔍 Fallback: dernière métrique en col {last_metric_col}")
            return last_metric_col
            
        except Exception as e:
            print(f"❌ Erreur: {e}")
            return 12
    
    # Test de la fonction
    end_col = find_convergence_table_end(ws)
    print(f"\n📊 L'extracteur va lire jusqu'à la colonne {end_col}")
    
    # Voir quels en-têtes il va lire
    print(f"\n📋 En-têtes extraits (colonnes 1 à {end_col}):")
    headers = {}
    excluded_headers = ['Date', 'Temps (j)', 'jours écoulés', 'Distance au front']
    
    for col_idx in range(1, end_col + 1):
        header_cell = ws.cell(9, col_idx).value
        if header_cell:
            header_text = str(header_cell).strip()
            if not any(excl.lower() in header_text.lower() for excl in excluded_headers):
                headers[col_idx] = header_text
                print(f"  Col {col_idx}: '{header_text}'")
    
    # Voir les vraies valeurs dans ces colonnes pour la ligne 66 (2025-08-25)
    print(f"\n📈 Valeurs finales (ligne 66, 2025-08-25):")
    for col_idx, header in headers.items():
        value = ws.cell(66, col_idx).value
        print(f"  {header}: {value}")
    
    # Comparer avec les colonnes du vrai tableau de convergences (4-9)
    print(f"\n📈 Vraies convergences (colonnes 4-9, ligne 66):")
    for col_idx in range(4, 10):
        header = ws.cell(9, col_idx).value
        value = ws.cell(66, col_idx).value
        print(f"  Col {col_idx} ({header}): {value}")

if __name__ == "__main__":
    test_extractor_columns()
