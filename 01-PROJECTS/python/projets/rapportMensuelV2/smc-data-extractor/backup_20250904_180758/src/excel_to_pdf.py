#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
SMC Excel to PDF - Impression des feuilles Excel
Convertit les feuilles Convergences et Déplacements des fichiers SMC en PDF
en utilisant les fonctions python
"""

import pandas as pd
import matplotlib.pyplot as plt
from pathlib import Path
from datetime import datetime
import numpy as np
from typing import Dict, List, Optional
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch, cm
from openpyxl import load_workbook
import matplotlib.dates as mdates

class SMCExcelToPDF:
    """Convertisseur Excel vers PDF pour les feuilles SMC"""
    
    def __init__(self, output_folder: str, month: str):
        self.output_folder = Path(output_folder)
        self.month = month
        
        # Styles pour PDF
        self.styles = getSampleStyleSheet()
        self.title_style = ParagraphStyle(
            'CustomTitle',
            parent=self.styles['Heading1'],
            fontSize=14,
            textColor=colors.darkblue,
            spaceAfter=12,
            alignment=1  # Center
        )
        
        self.heading_style = ParagraphStyle(
            'CustomHeading',
            parent=self.styles['Heading2'],
            fontSize=11,
            textColor=colors.darkgreen,
            spaceAfter=6
        )

    def extract_galerie_section(self, file_path: Path) -> tuple:
        """Extrait galerie et section depuis le nom de fichier"""
        file_name = file_path.stem.upper()
        
        # Patterns pour galerie (GGS, GVA, GHA, etc.)
        import re
        galerie_pattern = re.compile(r'([A-Z]{3})')
        galerie_matches = galerie_pattern.findall(file_name)
        galerie = galerie_matches[0] if galerie_matches else "UNKNOWN"
        
        # Pattern pour section (C003, T15, etc.)
        section_pattern = re.compile(r'([CT]\d{2,3})')
        section_matches = section_pattern.findall(file_name)
        section = section_matches[0] if section_matches else "UNKNOWN"
        
        return galerie, section

    def worksheet_to_table(self, worksheet, sheet_name: str, file_path: Path) -> List:
        """Convertit une feuille Excel en données pour tableau PDF"""
        
        # Extraire le titre et les infos du fichier
        galerie, section = self.extract_galerie_section(file_path)
        
        story_elements = []
        
        # Titre de la section
        title_text = f"{sheet_name} - {galerie} {section}"
        title = Paragraph(title_text, self.heading_style)
        story_elements.append(title)
        story_elements.append(Spacer(1, 0.1*inch))
        
        # Info fichier
        info_text = f"<i>Source: {file_path.name}</i>"
        info = Paragraph(info_text, self.styles['Normal'])
        story_elements.append(info)
        story_elements.append(Spacer(1, 0.1*inch))
        
        # Déterminer la zone de données selon la feuille
        if sheet_name == 'Convergences':
            # Pour Convergences: lignes 8-66, colonnes 1-9 (éviter les vitesses)
            start_row, end_row = 8, min(66, worksheet.max_row)
            start_col, end_col = 1, 9
        elif sheet_name == 'Déplacements':
            # Pour Déplacements: lignes 8-66, toutes les colonnes utiles
            start_row, end_row = 8, min(66, worksheet.max_row)
            start_col, end_col = 1, min(22, worksheet.max_column)
        else:
            return story_elements
        
        # Extraire les données
        table_data = []
        
        for row_idx in range(start_row, end_row + 1):
            row_data = []
            has_data = False
            
            for col_idx in range(start_col, end_col + 1):
                cell = worksheet.cell(row_idx, col_idx)
                cell_value = cell.value
                
                if cell_value is not None:
                    # Formatage des valeurs
                    if isinstance(cell_value, (int, float)):
                        if cell_value == int(cell_value):
                            formatted_value = str(int(cell_value))
                        else:
                            formatted_value = f"{cell_value:.3f}"
                    elif isinstance(cell_value, datetime):
                        formatted_value = cell_value.strftime('%Y-%m-%d')
                    else:
                        formatted_value = str(cell_value)
                    
                    row_data.append(formatted_value)
                    has_data = True
                else:
                    row_data.append("")
            
            # Ajouter la ligne si elle contient des données
            if has_data or row_idx <= start_row + 5:  # Toujours inclure les premières lignes
                table_data.append(row_data)
            
            # Arrêter si on a trop de lignes vides consécutives
            if not has_data and len(table_data) > 10:
                empty_rows = 0
                for check_row in table_data[-5:]:
                    if all(cell == "" for cell in check_row):
                        empty_rows += 1
                if empty_rows >= 5:
                    break
        
        if table_data:
            # Créer le tableau
            table = Table(table_data, repeatRows=1)
            
            # Style du tableau
            table_style = [
                ('FONTSIZE', (0, 0), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]
            
            # Style spécial pour les en-têtes (lignes 1-2)
            if len(table_data) > 1:
                table_style.extend([
                    ('BACKGROUND', (0, 0), (-1, 1), colors.lightblue),
                    ('FONTNAME', (0, 0), (-1, 1), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 1), 9),
                ])
            
            table.setStyle(TableStyle(table_style))
            story_elements.append(table)
        
        return story_elements

    def create_chart_from_excel(self, worksheet, sheet_name: str, file_path: Path) -> Optional[str]:
        """Crée un graphique simple depuis les données Excel"""
        try:
            galerie, section = self.extract_galerie_section(file_path)
            
            # Extraire les données pour le graphique
            dates = []
            values = {}
            
            # Lignes de données (à partir de la ligne 10)
            for row_idx in range(10, min(66, worksheet.max_row + 1)):
                # Date (colonne 2)
                date_cell = worksheet.cell(row_idx, 2).value
                if date_cell and isinstance(date_cell, datetime):
                    dates.append(date_cell)
                    
                    # Valeurs métriques
                    if sheet_name == 'Convergences':
                        # BG, HG, HD, BD (colonnes 4-7)
                        metrics = ['BG', 'HG', 'HD', 'BD']
                        for i, metric in enumerate(metrics, 4):
                            value = worksheet.cell(row_idx, i).value
                            if isinstance(value, (int, float)) and value != 0:
                                if metric not in values:
                                    values[metric] = []
                                values[metric].append(value)
                            else:
                                if metric not in values:
                                    values[metric] = []
                                values[metric].append(None)
                    
                    elif sheet_name == 'Déplacements':
                        # DPM colonnes (3, 6, 9, 12, 15)
                        metrics = ['DPM_BG', 'DPM_HG', 'DPM_CH', 'DPM_HD', 'DPM_BD']
                        cols = [3, 6, 9, 12, 15]
                        for metric, col in zip(metrics, cols):
                            value = worksheet.cell(row_idx, col).value
                            if isinstance(value, (int, float)) and value != 0:
                                if metric not in values:
                                    values[metric] = []
                                values[metric].append(value)
                            else:
                                if metric not in values:
                                    values[metric] = []
                                values[metric].append(None)
            
            if not dates or not values:
                return None
            
            # Créer le graphique
            plt.figure(figsize=(12, 8))
            
            for metric, metric_values in values.items():
                if any(v is not None for v in metric_values):
                    # Filtrer les valeurs None
                    clean_dates = []
                    clean_values = []
                    for d, v in zip(dates, metric_values):
                        if v is not None:
                            clean_dates.append(d)
                            clean_values.append(v)
                    
                    if clean_dates and clean_values:
                        plt.plot(clean_dates, clean_values, marker='o', label=metric, linewidth=2)
            
            plt.title(f'{sheet_name} - {galerie} {section} - {self.month}', fontsize=14, fontweight='bold')
            plt.xlabel('Date')
            plt.ylabel('Valeur (mm)')
            plt.legend()
            plt.grid(True, alpha=0.3)
            plt.xticks(rotation=45)
            plt.tight_layout()
            
            # Sauvegarder
            chart_path = self.output_folder / f"chart_{sheet_name}_{galerie}_{section}_{self.month.replace('-', '_')}.png"
            plt.savefig(chart_path, dpi=300, bbox_inches='tight')
            plt.close()
            
            return str(chart_path)
            
        except Exception as e:
            print(f"❌ Erreur création graphique {sheet_name}: {e}")
            return None

    def process_excel_file(self, file_path: Path) -> List:
        """Traite un fichier Excel et retourne les éléments pour le PDF"""
        story_elements = []
        
        try:
            workbook = load_workbook(file_path, data_only=True)
            galerie, section = self.extract_galerie_section(file_path)
            
            # Titre du fichier
            file_title = f"FICHIER: {galerie} {section}"
            title = Paragraph(file_title, self.title_style)
            story_elements.append(title)
            story_elements.append(Spacer(1, 0.2*inch))
            
            # Traiter les feuilles
            target_sheets = ['Convergences', 'Déplacements']
            
            for sheet_name in target_sheets:
                if sheet_name in workbook.sheetnames:
                    worksheet = workbook[sheet_name]
                    
                    # Ajouter le tableau
                    table_elements = self.worksheet_to_table(worksheet, sheet_name, file_path)
                    story_elements.extend(table_elements)
                    
                    # Ajouter le graphique
                    chart_path = self.create_chart_from_excel(worksheet, sheet_name, file_path)
                    if chart_path and Path(chart_path).exists():
                        story_elements.append(Spacer(1, 0.2*inch))
                        chart_title = Paragraph(f"Graphique {sheet_name}", self.styles['Heading3'])
                        story_elements.append(chart_title)
                        
                        from reportlab.platypus import Image
                        story_elements.append(Image(chart_path, width=10*inch, height=6*inch))
                    
                    story_elements.append(PageBreak())
                
        except Exception as e:
            print(f"❌ Erreur traitement {file_path}: {e}")
        
        return story_elements

def print_excel_sheets_to_pdf(root_folder: str, output_folder: str, month: str, log_callback=None) -> bool:
    """Fonction principale pour imprimer les feuilles Excel en PDF"""
    
    if log_callback:
        log_callback("📄 Impression des feuilles Excel en PDF...")
    
    try:
        # Créer le générateur
        generator = SMCExcelToPDF(output_folder, month)
        
        # Scanner les fichiers SMC
        root_path = Path(root_folder)
        excel_files = []
        
        for file_path in root_path.rglob("*.xlsm"):
            if 'SMC' in file_path.name.upper():
                excel_files.append(file_path)
        
        if not excel_files:
            if log_callback:
                log_callback("❌ Aucun fichier SMC trouvé")
            return False
        
        if log_callback:
            log_callback(f"📊 {len(excel_files)} fichiers SMC trouvés")
        
        # Créer le PDF avec le bon nom de fichier
        galerie, section = generator.extract_galerie_section(excel_files[0])
        pdf_filename = f"{excel_files[0].stem}-A4_landscape_marged.pdf"
        pdf_path = Path(output_folder) / pdf_filename
        
        doc = SimpleDocTemplate(str(pdf_path), pagesize=landscape(A4),
                              rightMargin=2*cm, leftMargin=2*cm,  # Marges comme référence
                              topMargin=2*cm, bottomMargin=2*cm)
        
        story = []
        
        # Titre principal
        main_title = Paragraph(f"Feuilles Excel SMC - {month}", generator.title_style)
        story.append(main_title)
        story.append(Spacer(1, 0.3*inch))
        
        # Traiter chaque fichier
        for i, file_path in enumerate(excel_files, 1):
            if log_callback:
                log_callback(f"📈 Traitement {i}/{len(excel_files)}: {file_path.name}")
            
            file_elements = generator.process_excel_file(file_path)
            story.extend(file_elements)
        
        # Générer le PDF
        if log_callback:
            log_callback("📄 Génération du PDF...")
        
        doc.build(story)
        
        if log_callback:
            log_callback(f"✅ PDF généré: {pdf_path}")
        
        return True
        
    except Exception as e:
        if log_callback:
            log_callback(f"❌ Erreur génération PDF: {str(e)}")
        return False

if __name__ == "__main__":
    # Test
    root_folder = r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250901-rapport mensuel aout\1-tableaux"
    output_folder = r"C:\temp\smc_output"
    month = "2025-08"
    
    print_excel_sheets_to_pdf(root_folder, output_folder, month)