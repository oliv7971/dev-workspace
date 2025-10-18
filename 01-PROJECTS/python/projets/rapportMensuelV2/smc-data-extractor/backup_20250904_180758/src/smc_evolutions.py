#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
SMC Data Extractor - Version FINALE
Extrait et calcule les évolutions de convergences et déplacements depuis les fichiers Excel SMC
"""

import os
import re
import logging
import pandas as pd
from pathlib import Path
from datetime import datetime, date
from typing import List, Dict, Any, Optional, Tuple
from openpyxl import load_workbook
from calendar import monthrange

# Configuration du logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class SMCExtractor:
    """Extracteur principal pour les données SMC"""
    
    def __init__(self, target_month: str):
        self.target_year, self.target_month = map(int, target_month.split('-'))
        
        # Fenêtre temporelle du mois
        self.month_start = date(self.target_year, self.target_month, 1)
        last_day = monthrange(self.target_year, self.target_month)[1]
        self.month_end = date(self.target_year, self.target_month, last_day)
        
        logger.info(f"Extracteur initialisé pour période: {self.month_start} à {self.month_end}")

    def parse_date(self, date_value) -> Optional[date]:
        """Parse une date dans différents formats"""
        if pd.isna(date_value) or date_value is None:
            return None
            
        # Si c'est déjà un datetime
        if isinstance(date_value, datetime):
            return date_value.date()
        
        # Si c'est un nombre Excel (serial date)
        if isinstance(date_value, (int, float)):
            try:
                # Excel epoch starts 1900-01-01 mais avec bug année bissextile 1900
                excel_epoch = datetime(1899, 12, 30)
                return (excel_epoch + pd.Timedelta(days=date_value)).date()
            except:
                return None
        
        # Si c'est une string
        if isinstance(date_value, str):
            # Nettoyer d'abord
            clean_date = date_value.strip()
            
            # Formats courants
            date_formats = [
                '%Y-%m-%d %H:%M:%S',  # 2024-01-15 00:00:00
                '%Y-%m-%d',
                '%d/%m/%Y',
                '%d/%m/%y',
                '%d-%m-%Y',
                '%d-%m-%y'
            ]
            
            for fmt in date_formats:
                try:
                    return datetime.strptime(clean_date, fmt).date()
                except:
                    continue
        
        return None

    def clean_numeric_value(self, value) -> Optional[float]:
        """Nettoie et convertit une valeur numérique"""
        if pd.isna(value) or value is None:
            return None
            
        if isinstance(value, (int, float)):
            return float(value)
        
        if isinstance(value, str):
            # Nettoyage des caractères spéciaux
            clean_value = value.replace(',', '.').replace(' ', '').replace('\u2212', '-')
            
            # Valeurs non-numériques
            if clean_value.upper() in ['ND', 'NR', '-', '']:
                return None
            
            try:
                return float(clean_value)
            except:
                return None
        
        return None

    def extract_sheet_data(self, worksheet, sheet_name: str, file_path: Path) -> Dict[str, Any]:
        """Extrait les données d'une feuille selon la structure SMC"""
        
        if sheet_name == 'Convergences':
            # Pour CONVERGENCES: uniquement le 1er tableau (VALEURS) - colonnes 1 à 9
            headers_row = 9
            data_start_row = 10
            
            # Extraction des en-têtes SEULEMENT jusqu'à la colonne 9 (LB)
            headers = {}
            for col_idx in range(1, 10):  # Colonnes 1 à 9 seulement
                header_cell = worksheet.cell(headers_row, col_idx).value
                if header_cell:
                    headers[col_idx] = str(header_cell).strip()
            
        elif sheet_name == 'Déplacements':
            # Pour DÉPLACEMENTS: structure normale - toutes les colonnes
            headers_row = 9
            data_start_row = 10
            
            # Extraction de tous les en-têtes
            headers = {}
            for col_idx in range(1, worksheet.max_column + 1):
                header_cell = worksheet.cell(headers_row, col_idx).value
                if header_cell:
                    headers[col_idx] = str(header_cell).strip()
        
        else:
            return {}
        
        # Trouver la colonne de date (toujours colonne 2)
        date_col = 2
        
        # Extraction des données
        data_rows = []
        
        for row_idx in range(data_start_row, worksheet.max_row + 1):
            # Date
            date_cell = worksheet.cell(row_idx, date_col).value
            parsed_date = self.parse_date(date_cell)
            if not parsed_date:
                continue
            
            row_data = {'date': parsed_date}
            
            # Autres colonnes (selon les en-têtes extraits)
            for col_idx, header in headers.items():
                if col_idx == date_col:
                    continue  # Date déjà traitée
                
                cell_value = worksheet.cell(row_idx, col_idx).value
                clean_value = self.clean_numeric_value(cell_value)
                
                if clean_value is not None:
                    row_data[header] = clean_value
            
            if len(row_data) > 1:  # Au moins date + 1 métrique
                data_rows.append(row_data)
        
        return {
            'sheet_name': sheet_name,
            'file_path': str(file_path),
            'data': data_rows,
            'headers': headers
        }

    def extract_galerie_section(self, file_path: Path) -> Tuple[str, str]:
        """Extrait galerie et section depuis le nom de fichier"""
        file_name = file_path.stem.upper()
        
        # Patterns pour galerie (GGS, GVA, GHA, etc.)
        galerie_pattern = re.compile(r'([A-Z]{3})')
        galerie_matches = galerie_pattern.findall(file_name)
        galerie = galerie_matches[0] if galerie_matches else "UNKNOWN"
        
        # Pattern pour section (C003, T15, etc.)
        section_pattern = re.compile(r'([CT]\d{2,3})')
        section_matches = section_pattern.findall(file_name)
        section = section_matches[0] if section_matches else "UNKNOWN"
        
        return galerie, section

    def calculate_evolutions(self, data_rows: List[Dict], metrics: List[str]) -> Dict[str, Dict]:
        """Calcule les évolutions mensuelles et cumulées"""
        if not data_rows:
            return {}
        
        # Filtrer par période
        dates_in_month = []
        dates_before_month = []
        
        for row in data_rows:
            row_date = row['date']
            if self.month_start <= row_date <= self.month_end:
                dates_in_month.append(row)
            elif row_date < self.month_start:
                dates_before_month.append(row)
        
        # Trier par date
        dates_in_month.sort(key=lambda x: x['date'])
        dates_before_month.sort(key=lambda x: x['date'])
        
        # TOUTES LES DONNÉES triées par date (pour valeur cumulée)
        all_data_sorted = sorted(data_rows, key=lambda x: x['date'])
        
        results = {}
        
        for metric in metrics:
            if metric not in ['Date', 'Temps (j)', 'jours écoulés', 'Distance au front']:
                # Dernière valeur dans le mois
                last_in_month_value = None
                last_in_month_date = None
                for row in reversed(dates_in_month):
                    if metric in row and row[metric] is not None:
                        last_in_month_value = row[metric]
                        last_in_month_date = row['date']
                        break
                
                # Dernière valeur avant le mois
                last_before_month_value = None
                last_before_month_date = None
                for row in reversed(dates_before_month):
                    if metric in row and row[metric] is not None:
                        last_before_month_value = row[metric]
                        last_before_month_date = row['date']
                        break
                
                # DERNIÈRE VALEUR CUMULÉE (de toutes les données)
                last_cumulative_value = None
                last_cumulative_date = None
                for row in reversed(all_data_sorted):
                    if metric in row and row[metric] is not None:
                        last_cumulative_value = row[metric]
                        last_cumulative_date = row['date']
                        break
                
                # Calculs
                periodic_mm = None
                if last_in_month_value is not None and last_before_month_value is not None:
                    periodic_mm = round(last_in_month_value - last_before_month_value, 2)
                
                cumulative_mm = None
                if last_cumulative_value is not None:
                    cumulative_mm = round(last_cumulative_value, 2)
                
                results[metric] = {
                    'metric': metric,
                    'periodic_mm': periodic_mm,
                    'cumulative_mm': cumulative_mm,
                    'date_last_in_month': last_in_month_date,
                    'date_last_before_month': last_before_month_date,
                    'date_last_cumulative': last_cumulative_date,
                    'nb_dates_in_month': len(dates_in_month),
                    'nb_dates_before_month': len(dates_before_month)
                }
        
        return results

    def process_file(self, file_path: Path) -> Dict[str, Any]:
        """Traite un fichier Excel"""
        
        try:
            workbook = load_workbook(file_path, data_only=True)
            
            file_results = {
                'file_path': str(file_path),
                'convergences': [],
                'deplacements': [],
                'dates': []
            }
            
            # Extraire galerie et section
            galerie, section = self.extract_galerie_section(file_path)
            
            # Traiter les feuilles
            target_sheets = {
                'Convergences': 'convergences',
                'Déplacements': 'deplacements'
            }
            
            for sheet_name, result_key in target_sheets.items():
                if sheet_name in workbook.sheetnames:
                    worksheet = workbook[sheet_name]
                    sheet_data = self.extract_sheet_data(worksheet, sheet_name, file_path)
                    
                    if sheet_data and sheet_data['data']:
                        # Calculer évolutions
                        metrics = list(sheet_data['headers'].values())
                        evolutions = self.calculate_evolutions(sheet_data['data'], metrics)
                        
                        # Formatter les résultats
                        for metric, evolution in evolutions.items():
                            result = {
                                'source_file': str(file_path),
                                'sheet': sheet_name,
                                'galerie': galerie,
                                'section': section,
                                'month': f"{self.target_year}-{self.target_month:02d}",
                                **evolution
                            }
                            file_results[result_key].append(result)
                        
                        # Dates du mois
                        dates_in_month = [
                            row['date'] for row in sheet_data['data'] 
                            if self.month_start <= row['date'] <= self.month_end
                        ]
                        
                        if dates_in_month:
                            date_result = {
                                'source_file': str(file_path),
                                'sheet': sheet_name,
                                'galerie': galerie,
                                'section': section,
                                'month': f"{self.target_year}-{self.target_month:02d}",
                                'dates_in_month': [d.strftime('%Y-%m-%d') for d in sorted(set(dates_in_month))],
                                'total_dates_in_month': len(set(dates_in_month)),
                                'first_date_in_month': min(dates_in_month).strftime('%Y-%m-%d'),
                                'last_date_in_month': max(dates_in_month).strftime('%Y-%m-%d')
                            }
                            file_results['dates'].append(date_result)
            
            return file_results
            
        except Exception as e:
            logger.error(f"Erreur lors du traitement de {file_path}: {str(e)}")
            return {}

def run_extraction(root_folder: str, month: str, output_folder: str, options: Dict = None, log_callback=None):
    """Point d'entrée principal pour l'extraction FINALE"""
    if not options:
        options = {'generate_csv': True, 'generate_pdf': False}
    
    if log_callback:
        log_callback(f"🚀 EXTRACTION FINALE SMC - Mois: {month}")
        log_callback(f"📁 Source: {root_folder}")
        log_callback(f"💾 Sortie: {output_folder}")
    
    # Validation du mois
    try:
        datetime.strptime(month, '%Y-%m')
    except ValueError:
        if log_callback:
            log_callback(f"❌ Format mois invalide: {month}")
        return False
    
    # Créer le dossier de sortie
    output_path = Path(output_folder)
    output_path.mkdir(parents=True, exist_ok=True)
    
    # Initialiser l'extracteur
    extractor = SMCExtractor(month)
    
    # Scanner les fichiers
    if log_callback:
        log_callback("🔍 Recherche des fichiers SMC...")
    
    root_path = Path(root_folder)
    excel_files = []
    
    for file_path in root_path.rglob("*.xlsm"):
        if 'SMC' in file_path.name.upper():
            excel_files.append(file_path)
    
    if not excel_files:
        if log_callback:
            log_callback("❌ Aucun fichier SMC trouvé")
        return False
    
    if log_callback:
        log_callback(f"📊 {len(excel_files)} fichiers SMC trouvés")
    
    # Traiter chaque fichier
    all_convergences = []
    all_deplacements = []
    all_dates = []
    
    for i, file_path in enumerate(excel_files, 1):
        if log_callback:
            log_callback(f"📈 Traitement {i}/{len(excel_files)}: {file_path.name}")
        
        file_results = extractor.process_file(file_path)
        if file_results:
            all_convergences.extend(file_results.get('convergences', []))
            all_deplacements.extend(file_results.get('deplacements', []))
            all_dates.extend(file_results.get('dates', []))
            
            if log_callback:
                conv_count = len(file_results.get('convergences', []))
                depl_count = len(file_results.get('deplacements', []))
                log_callback(f"   ✅ {conv_count} convergences, {depl_count} déplacements")
    
    # Exporter les résultats en CSV
    if options.get('generate_csv', True):
        if log_callback:
            log_callback("")
            log_callback("💾 Génération des fichiers CSV...")
        
        # Export des dates
        if all_dates:
            dates_df = pd.DataFrame(all_dates)
            dates_path = output_path / "Dates_SMC.csv"
            dates_df.to_csv(dates_path, index=False, encoding='utf-8-sig')
            if log_callback:
                log_callback(f"✅ Dates_SMC.csv: {len(all_dates)} entrées")
        
        # Export des convergences
        if all_convergences:
            conv_df = pd.DataFrame(all_convergences)
            conv_path = output_path / "Convergences_SMC.csv"
            conv_df.to_csv(conv_path, index=False, encoding='utf-8-sig')
            if log_callback:
                log_callback(f"✅ Convergences_SMC.csv: {len(all_convergences)} métriques")
        
        # Export des déplacements
        if all_deplacements:
            depl_df = pd.DataFrame(all_deplacements)
            depl_path = output_path / "Deplacements_SMC.csv"
            depl_df.to_csv(depl_path, index=False, encoding='utf-8-sig')
            if log_callback:
                log_callback(f"✅ Deplacements_SMC.csv: {len(all_deplacements)} métriques")
    
    # Générer les récapitulatifs en ligne
    if options.get('generate_ligne_summaries', True):
        if log_callback:
            log_callback("")
            log_callback("📋 Génération des récapitulatifs en ligne...")
        
        try:
            from ligne_summary_generator import generate_ligne_summaries
            ligne_success = generate_ligne_summaries(str(output_path), str(output_path), month, log_callback)
            
            if not ligne_success and log_callback:
                log_callback("⚠️ Génération récapitulatifs en ligne échouée")
                
        except ImportError as e:
            if log_callback:
                log_callback(f"⚠️ Module récapitulatifs en ligne manquant: {e}")
        except Exception as e:
            if log_callback:
                log_callback(f"⚠️ Erreur récapitulatifs en ligne: {e}")

    # Générer le PDF si demandé
    if options.get('generate_pdf', False):
        if log_callback:
            log_callback("")
            log_callback("📄 Génération des PDF...")
        
        try:
            # Impression native Excel (format original)
            from excel_to_pdf_native import print_excel_files_native
            native_pdf_success = print_excel_files_native(root_folder, str(output_path), month, log_callback)
            
            # Optionnel: PDF de synthèse avec graphiques Python
            # from excel_to_pdf import print_excel_sheets_to_pdf
            # python_pdf_success = print_excel_sheets_to_pdf(root_folder, str(output_path), month, log_callback)
            
            if not native_pdf_success and log_callback:
                log_callback("⚠️ Impression native échouée")
                
        except ImportError as e:
            if log_callback:
                log_callback(f"⚠️ Module impression native manquant: {e}")
        except Exception as e:
            if log_callback:
                log_callback(f"⚠️ Erreur impression: {e}")
    
    if log_callback:
        log_callback("")
        log_callback(f"🎉 EXTRACTION TERMINÉE AVEC SUCCÈS!")
        log_callback(f"📊 Total: {len(all_dates)} feuilles, {len(all_convergences)} convergences, {len(all_deplacements)} déplacements")
        log_callback(f"📁 Fichiers dans: {output_folder}")
    
    return True

if __name__ == "__main__":
    # Test avec des valeurs par défaut
    root_folder = r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250901-rapport mensuel aout\1-tableaux"
    month = "2025-08"
    output_folder = r"C:\temp\smc_output"
    
    run_extraction(root_folder, month, output_folder)