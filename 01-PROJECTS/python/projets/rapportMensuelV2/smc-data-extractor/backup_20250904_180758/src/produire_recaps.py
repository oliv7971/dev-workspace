import os
import pandas as pd
from docx import Document

def integrate_csv_to_word(csv_files, output_docx):
    document = Document()
    document.add_heading('Rapport d\'intégration des CSV', level=1)

    for csv_file in csv_files:
        if os.path.exists(csv_file):
            document.add_heading(os.path.basename(csv_file), level=2)
            df = pd.read_csv(csv_file)
            table = document.add_table(rows=1, cols=len(df.columns))
            hdr_cells = table.rows[0].cells
            for i, column_name in enumerate(df.columns):
                hdr_cells[i].text = column_name

            for index, row in df.iterrows():
                row_cells = table.add_row().cells
                for i, value in enumerate(row):
                    row_cells[i].text = str(value)

            document.add_page_break()
        else:
            print(f"Le fichier {csv_file} n'existe pas et ne peut pas être intégré.")

    document.save(output_docx)

if __name__ == "__main__":
    # Exemple d'utilisation
    csv_files = [
        'path/to/Dates.csv',
        'path/to/Convergences.csv',
        'path/to/Deplacements.csv'
    ]
    output_docx = 'path/to/output_report.docx'
    integrate_csv_to_word(csv_files, output_docx)