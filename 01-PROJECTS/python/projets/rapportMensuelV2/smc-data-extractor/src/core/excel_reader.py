from openpyxl import load_workbook
import os
import re

class ExcelReader:
    def __init__(self, file_path):
        self.file_path = file_path
        self.workbook = load_workbook(filename=file_path, data_only=True)
        self.data = {}

    def extract_data(self):
        for sheet_name in self.workbook.sheetnames:
            sheet = self.workbook[sheet_name]
            self.data[sheet_name] = self._parse_sheet(sheet)

    def _parse_sheet(self, sheet):
        parsed_data = {
            'convergences': [],
            'deplacements': [],
            'dates': []
        }
        
        date_column_index = self._find_date_column(sheet)
        if date_column_index is None:
            return parsed_data

        for row in sheet.iter_rows(min_row=3, values_only=True):
            date_value = row[date_column_index]
            if date_value:
                parsed_data['dates'].append(self._normalize_date(date_value))
                parsed_data['convergences'].append(self._extract_convergences(row))
                parsed_data['deplacements'].append(self._extract_deplacements(row))

        return parsed_data

    def _find_date_column(self, sheet):
        for row in sheet.iter_rows(min_row=1, max_row=3, values_only=True):
            for index, cell in enumerate(row):
                if self._is_date_header(cell):
                    return index
        return None

    def _is_date_header(self, cell_value):
        return isinstance(cell_value, str) and re.search(r'(?i)date|jour|heure', cell_value)

    def _normalize_date(self, date_value):
        if isinstance(date_value, str):
            # Implement normalization logic for string dates
            return date_value
        return date_value.isoformat()  # Assuming date_value is a datetime object

    def _extract_convergences(self, row):
        # Implement logic to extract convergence data from the row
        return {}

    def _extract_deplacements(self, row):
        # Implement logic to extract displacement data from the row
        return {}

    def get_data(self):
        return self.data

def main():
    folder_path = 'C:\\data\\11-CHANTIERS\\BURE\\10-ACTIVITES\\Rapports d\'activité\\2025\\250901-rapport mensuel aout\\1-tableaux'
    for filename in os.listdir(folder_path):
        if filename.endswith('.xlsm') or filename.endswith('.xlsx'):
            file_path = os.path.join(folder_path, filename)
            reader = ExcelReader(file_path)
            reader.extract_data()
            data = reader.get_data()
            print(f'Data extracted from {filename}: {data}')

if __name__ == '__main__':
    main()