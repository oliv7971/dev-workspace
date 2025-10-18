import os
import glob
import pandas as pd
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

# === CONFIGURATION ===
DEBUG = False  # 🔁 Mets True pour tester structure, False pour générer le fichier final
#input_folder = r"C:\data\11-CHANTIERS\FREJUS\2025-07-14-traitement FRejus 2025\02-conversion xlsx\01-PARTIE FRANCE"
input_folder = r"C:\data\11-CHANTIERS\FREJUS\2025-07-14-traitement FRejus 2025\02-conversion xlsx\02-PARTIE ITALIE"
output_file = "synthese_convergences_script_final.xlsx"

xlsx_files = sorted(glob.glob(os.path.join(input_folder, "*.xlsx")))
all_data = []

for file in xlsx_files:
    print(f"\n🔎 Traitement de : {os.path.basename(file)}")

    try:
        wb = openpyxl.load_workbook(file, data_only=True)
        if "tableau" not in wb.sheetnames:
            print("❌ Onglet 'tableau' introuvable.")
            continue

        ws = wb["tableau"]
        df_temp = pd.DataFrame(ws.values)

        header_row_index = None
        for i, row in df_temp.iterrows():
            if isinstance(row[0], str) and "DATE" in row[0].upper():
                header_row_index = i
                break
        if header_row_index is None:
            print("❌ Ligne contenant 'DATE' introuvable.")
            continue

        base_cols = {}
        for col_idx, cell in enumerate(df_temp.iloc[header_row_index - 2]):
            if isinstance(cell, str):
                for base in ["base 1-4", "base 2-3"]:
                    if base in cell.lower():
                        base_cols[base] = col_idx

        if not base_cols:
            print("❌ Colonnes 'base 1-4' ou 'base 2-3' introuvables.")
            continue

        print(f"✅ Ok : colonnes détectées = {base_cols}")

        if DEBUG:
            continue

      # Extraction des données
        rows = []
        for i in range(header_row_index + 1, len(df_temp)):
            row_data = {
                "nom_fichier": os.path.basename(file),
                "Date": df_temp.iat[i, 0]
            }
            for base, idx in base_cols.items():
                try:
                    long_val = pd.to_numeric(df_temp.iat[i, idx], errors='coerce')
                    conv_val = pd.to_numeric(df_temp.iat[i, idx + 1], errors='coerce')
                except Exception:
                    long_val = conv_val = None
                row_data[f"LONG {base[-3:]}"] = long_val
                row_data[f"CONV {base[-3:]}"] = conv_val
            rows.append(row_data)

        df_rows = pd.DataFrame(rows)

        # ✅ Nettoyage des lignes vides ou parasites
        df_rows = df_rows.dropna(how="all")  # Supprime toute ligne complètement vide
        df_rows = df_rows.dropna(subset=["Date"], how="all")  # Supprime les lignes sans date (souvent vides)

        df_rows["MOY CONV (1-4/2-3)"] = df_rows[["CONV 1-4", "CONV 2-3"]].mean(axis=1)
        all_data.append(df_rows)


    except Exception as e:
        print(f"❌ Erreur inattendue : {e}")

# === Fusion et export Excel ===
if not DEBUG:
    if not all_data:
        print("\n⚠️ Aucun fichier traité. Vérifie le mode DEBUG ou la structure des fichiers.")
        raise ValueError("Aucun fichier valide trouvé.")

    df_final_all = pd.concat(all_data, ignore_index=True)

    # Mise en forme Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"

    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(bottom=Side(style='thin'))
    thick_border = Border(
        top=Side(style="medium"),
        bottom=Side(style="medium"),
        left=Side(style="medium"),
        right=Side(style="medium")
    )
    number_format = "0.0"

    current_row = 1
    grouped = df_final_all.groupby("nom_fichier")
    col_count = len(df_final_all.columns) - 1  # sans nom_fichier

    for file_name, group_df in grouped:
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_count)
        title_cell = ws.cell(row=current_row, column=1, value=file_name)
        title_cell.alignment = center_alignment
        current_row += 1

        data_start_row = current_row
        data_end_row = current_row + len(group_df)

        for r_idx, row in enumerate(dataframe_to_rows(group_df.drop(columns="nom_fichier"), index=False, header=True), start=current_row):
            for c_idx, value in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.alignment = center_alignment
                if isinstance(value, (int, float)) and r_idx != current_row:
                    cell.number_format = number_format
            if r_idx == current_row:
                for col in range(1, len(row) + 1):
                    ws.cell(row=r_idx, column=col).border = thin_border

        for row in range(data_start_row, data_end_row + 1):
            for col in range(1, col_count + 1):
                cell = ws.cell(row=row, column=col)
                border = Border(
                    top=thick_border.top if row == data_start_row else None,
                    bottom=thick_border.bottom if row == data_end_row else None,
                    left=thick_border.left if col == 1 else None,
                    right=thick_border.right if col == col_count else None
                )
                cell.border = border

        current_row = data_end_row + 3  # Saut de lignes entre blocs

    for col_idx in range(1, col_count + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 15

    wb.save(output_file)
    print(f"\n✅ Fichier généré : {output_file}")
