"""
Générateur de fiches d'implantation à partir de données Leica
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import re


def parser_fichier_leica(fichier_txt):
    """
    Parse le fichier texte Leica et extrait les données des points
    """
    points = []
    
    with open(fichier_txt, 'r', encoding='utf-8', errors='ignore') as f:
        lignes = f.readlines()
    
    # Trouver la section "Carnet brut"
    debut_carnet = False
    for i, ligne in enumerate(lignes):
        if 'Carnet brut' in ligne:
            debut_carnet = True
            # Sauter les lignes d'en-tête
            continue
        
        if debut_carnet and ligne.strip() and not ligne.startswith('N°'):
            # Parser les données (format tabulé)
            parties = ligne.split('\t')
            if len(parties) >= 4:
                try:
                    nom = parties[0].strip()
                    est = parties[1].strip()
                    nord = parties[2].strip()
                    altitude = parties[3].strip()
                    
                    # Vérifier si on a des valeurs numériques
                    if est and nord and altitude:
                        try:
                            est_val = float(est)
                            nord_val = float(nord)
                            alt_val = float(altitude)
                            
                            # Ignorer les lignes de station
                            if not nom.startswith('System') and nom:
                                points.append({
                                    'Nom': nom,
                                    'Est': est_val,
                                    'Nord': nord_val,
                                    'Altitude': alt_val
                                })
                        except ValueError:
                            pass
                except IndexError:
                    pass
        
        # Arrêter à la fin du carnet brut
        if debut_carnet and 'System 1200 Fichier Journal' in ligne:
            break
    
    return points


def creer_fiche_implantation(points, fichier_sortie):
    """
    Crée un fichier Excel avec les fiches d'implantation
    """
    # Créer un nouveau classeur
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fiches d'implantation"
    
    # Styles
    titre_font = Font(name='Arial', size=14, bold=True)
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    data_font = Font(name='Arial', size=10)
    
    border_thin = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # En-tête du document
    ws.merge_cells('A1:F1')
    ws['A1'] = 'FICHES D\'IMPLANTATION'
    ws['A1'].font = titre_font
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A2:F2')
    ws['A2'] = f'Projet GMA - Date: {datetime.now().strftime("%d/%m/%Y")}'
    ws['A2'].font = Font(name='Arial', size=11)
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    
    # En-tête du tableau
    ligne = 4
    headers = ['N°', 'Nom du point', 'Est (m)', 'Nord (m)', 'Altitude (m)', 'Observations']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=ligne, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border_thin
    
    # Données des points
    ligne = 5
    for idx, point in enumerate(points, start=1):
        ws.cell(row=ligne, column=1, value=idx)
        ws.cell(row=ligne, column=2, value=point['Nom'])
        ws.cell(row=ligne, column=3, value=f"{point['Est']:.3f}")
        ws.cell(row=ligne, column=4, value=f"{point['Nord']:.3f}")
        ws.cell(row=ligne, column=5, value=f"{point['Altitude']:.3f}")
        ws.cell(row=ligne, column=6, value='')
        
        # Appliquer le style
        for col in range(1, 7):
            cell = ws.cell(row=ligne, column=col)
            cell.font = data_font
            cell.border = border_thin
            if col in [3, 4, 5]:  # Colonnes numériques
                cell.alignment = Alignment(horizontal='right', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        ligne += 1
    
    # Ajuster la largeur des colonnes
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 15
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 25
    
    # Sauvegarder le fichier
    wb.save(fichier_sortie)
    print(f"Fichier créé avec succès : {fichier_sortie}")
    print(f"Nombre de points traités : {len(points)}")


def main():
    """
    Fonction principale
    """
    fichier_entree = '20251203_GMA.txt'
    fichier_sortie = 'Fiches_implantation_GMA.xlsx'
    
    print("Lecture du fichier Leica...")
    points = parser_fichier_leica(fichier_entree)
    
    if not points:
        print("Aucun point trouvé dans le fichier.")
        return
    
    print(f"Points trouvés : {len(points)}")
    for point in points[:5]:  # Afficher les 5 premiers points
        print(f"  - {point['Nom']}: E={point['Est']:.3f}, N={point['Nord']:.3f}, Z={point['Altitude']:.3f}")
    
    if len(points) > 5:
        print(f"  ... et {len(points) - 5} autres points")
    
    print("\nCréation du fichier Excel...")
    creer_fiche_implantation(points, fichier_sortie)


if __name__ == '__main__':
    main()
