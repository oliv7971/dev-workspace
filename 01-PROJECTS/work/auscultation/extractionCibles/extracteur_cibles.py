#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extracteur de coordonnées de cibles depuis fichiers Excel
Extrait les coordonnées X, Y, Z de la dernière mesure de chaque cible

Auteur: Assistant IA
Date: Novembre 2025
"""

import openpyxl
import pandas as pd
import os
from pathlib import Path

class ExtracteurCibles:
    """
    Classe pour extraire les coordonnées des cibles depuis un fichier Excel
    """
    
    def __init__(self):
        self.fichier = None
        self.workbook = None
        self.worksheet = None
        self.cibles_extraites = []
    
    def charger_fichier(self, chemin_fichier):
        """
        Charge un fichier Excel
        
        Args:
            chemin_fichier (str): Chemin vers le fichier Excel
            
        Returns:
            bool: True si le chargement a réussi
        """
        try:
            self.fichier = chemin_fichier
            self.workbook = openpyxl.load_workbook(chemin_fichier, data_only=True)
            
            print(f"Fichier chargé: {os.path.basename(chemin_fichier)}")
            print(f"Onglets disponibles: {self.workbook.sheetnames}")
            
            # Chercher l'onglet "Résultats observations"
            if "Résultats observations" in self.workbook.sheetnames:
                self.worksheet = self.workbook["Résultats observations"]
                print(f"Onglet 'Résultats observations' sélectionné")
                return True
            else:
                print("ERREUR: Onglet 'Résultats observations' non trouvé")
                return False
                
        except Exception as e:
            print(f"Erreur lors du chargement: {e}")
            return False
    
    def detecter_structure(self):
        """
        Détecte automatiquement la structure du fichier
        (colonnes D à AA ou E à Y)
        
        Returns:
            tuple: (colonne_debut, colonne_fin, type_fichier)
        """
        # Vérifier si colonne D a un nom de cible
        col_d_value = self.worksheet.cell(row=8, column=4).value
        
        if col_d_value and col_d_value not in ['Date', 'Jours', 'Front', 'Dist']:
            # Type 1: colonnes D à AA
            return (4, 27, "D-AA")  # D=4, AA=27
        else:
            # Type 2: colonnes E à Y
            return (5, 25, "E-Y")   # E=5, Y=25
    
    def trouver_derniere_ligne_donnees(self, col_debut, col_fin):
        """
        Trouve la dernière ligne contenant des données
        
        Args:
            col_debut (int): Colonne de début
            col_fin (int): Colonne de fin
            
        Returns:
            int: Numéro de la dernière ligne avec données
        """
        for row in range(self.worksheet.max_row, 9, -1):
            # Vérifier s'il y a des données dans cette ligne
            for col in range(col_debut, col_fin + 1):
                val = self.worksheet.cell(row=row, column=col).value
                if val is not None and val != '':
                    return row
        return None
    
    def extraire_noms_cibles(self, col_debut, col_fin):
        """
        Extrait les noms des cibles depuis la ligne 8 (cellules fusionnées)
        
        Args:
            col_debut (int): Colonne de début
            col_fin (int): Colonne de fin
            
        Returns:
            list: Liste de tuples (nom_cible, col_debut_cible, col_fin_cible)
        """
        cibles = []
        
        # Parcourir toutes les fusions de la ligne 8
        for merged_cell in self.worksheet.merged_cells.ranges:
            if merged_cell.min_row == 8 and merged_cell.max_row == 8:
                start_col = merged_cell.min_col
                end_col = merged_cell.max_col
                
                # Vérifier si cette fusion est dans la plage demandée
                if start_col >= col_debut and end_col <= col_fin:
                    nom_cible = self.worksheet.cell(row=8, column=start_col).value
                    
                    # Vérifier que c'est bien une cible (pas "Différences" par exemple)
                    if nom_cible and nom_cible != "Différences":
                        # Vérifier qu'il y a bien 3 colonnes (X, Y, Z)
                        if end_col - start_col + 1 == 3:
                            cibles.append((nom_cible, start_col, end_col))
        
        # Trier par colonne
        cibles.sort(key=lambda x: x[1])
        
        return cibles
    
    def extraire_coordonnees(self):
        """
        Extrait les coordonnées X, Y, Z de toutes les cibles
        
        Returns:
            bool: True si l'extraction a réussi
        """
        if not self.worksheet:
            print("ERREUR: Aucun worksheet chargé")
            return False
        
        try:
            # Détecter la structure
            col_debut, col_fin, type_fichier = self.detecter_structure()
            print(f"\nType de fichier détecté: {type_fichier}")
            print(f"Colonnes: {openpyxl.utils.get_column_letter(col_debut)} à {openpyxl.utils.get_column_letter(col_fin)}")
            
            # Trouver la dernière ligne avec données
            derniere_ligne = self.trouver_derniere_ligne_donnees(col_debut, col_fin)
            if not derniere_ligne:
                print("ERREUR: Aucune donnée trouvée")
                return False
            
            print(f"Dernière ligne avec données: {derniere_ligne}")
            
            # Extraire les noms des cibles
            cibles_info = self.extraire_noms_cibles(col_debut, col_fin)
            print(f"Nombre de cibles trouvées: {len(cibles_info)}")
            
            # Extraire les coordonnées de chaque cible
            self.cibles_extraites = []
            
            for nom_cible, col_x, col_z in cibles_info:
                x = self.worksheet.cell(row=derniere_ligne, column=col_x).value
                y = self.worksheet.cell(row=derniere_ligne, column=col_x + 1).value
                z = self.worksheet.cell(row=derniere_ligne, column=col_x + 2).value
                
                # Ajouter seulement si au moins une coordonnée est présente
                if x is not None or y is not None or z is not None:
                    self.cibles_extraites.append({
                        'Cible': nom_cible,
                        'X': x if x is not None else '',
                        'Y': y if y is not None else '',
                        'Z': z if z is not None else ''
                    })
            
            print(f"Coordonnées extraites: {len(self.cibles_extraites)} cibles")
            return True
            
        except Exception as e:
            print(f"Erreur lors de l'extraction: {e}")
            import traceback
            traceback.print_exc()
            return False
    
    def afficher_resultats(self):
        """
        Affiche les résultats extraits
        """
        if not self.cibles_extraites:
            print("Aucune donnée extraite")
            return
        
        print(f"\n{'='*80}")
        print(f"{'CIBLE':<30} {'X':>15} {'Y':>15} {'Z':>15}")
        print(f"{'='*80}")
        
        for cible in self.cibles_extraites:
            x_str = f"{cible['X']:.4f}" if isinstance(cible['X'], (int, float)) else str(cible['X'])
            y_str = f"{cible['Y']:.4f}" if isinstance(cible['Y'], (int, float)) else str(cible['Y'])
            z_str = f"{cible['Z']:.4f}" if isinstance(cible['Z'], (int, float)) else str(cible['Z'])
            
            print(f"{cible['Cible']:<30} {x_str:>15} {y_str:>15} {z_str:>15}")
    
    def sauvegarder_csv(self, chemin_sortie=None):
        """
        Sauvegarde les résultats dans un fichier CSV
        
        Args:
            chemin_sortie (str): Chemin du fichier de sortie
            
        Returns:
            str: Chemin du fichier sauvegardé
        """
        if not self.cibles_extraites:
            print("Aucune donnée à sauvegarder")
            return None
        
        if chemin_sortie is None:
            # Générer un nom automatique
            nom_base = Path(self.fichier).stem
            chemin_sortie = f"{nom_base}_cibles_extraites.csv"
        
        try:
            df = pd.DataFrame(self.cibles_extraites)
            df.to_csv(chemin_sortie, sep=';', index=False, encoding='utf-8-sig')
            print(f"\nRésultats sauvegardés: {chemin_sortie}")
            return chemin_sortie
            
        except Exception as e:
            print(f"Erreur lors de la sauvegarde: {e}")
            return None


def main():
    """
    Fonction principale pour tester l'extracteur
    """
    print("="*80)
    print("EXTRACTEUR DE COORDONNÉES DE CIBLES")
    print("="*80)
    
    # Fichier exemple
    fichier = "01-GVA-PM 169_Tableau_Jour-22-12-20b.xlsx"
    
    if not os.path.exists(fichier):
        print(f"Fichier non trouvé: {fichier}")
        return
    
    # Créer l'extracteur
    extracteur = ExtracteurCibles()
    
    # Charger le fichier
    if not extracteur.charger_fichier(fichier):
        return
    
    # Extraire les coordonnées
    if not extracteur.extraire_coordonnees():
        return
    
    # Afficher les résultats
    extracteur.afficher_resultats()
    
    # Sauvegarder
    extracteur.sauvegarder_csv()
    
    print("\n" + "="*80)
    print("TRAITEMENT TERMINÉ")
    print("="*80)


if __name__ == "__main__":
    main()
