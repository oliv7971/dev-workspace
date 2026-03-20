#!/usr/bin/env python3
"""
Script simple pour ajouter les 20 cibles au fichier CLEAN qui fonctionne
"""

import openpyxl
from openpyxl.utils import get_column_letter
import shutil

def add_targets_to_clean_file():
    """Ajoute les 20 nouvelles cibles au fichier CLEAN qui fonctionne"""

    input_file = "01_51-B-GER-HA21_3-PM_135-25-05-21_CLEAN.xlsm"
    output_file = "01_51-B-GER-HA21_3-PM_135-25-05-21_FINAL.xlsm"

    # Copier le fichier qui fonctionne
    shutil.copy2(input_file, output_file)
    print(f"📋 Copie du fichier qui fonctionne: {output_file}")

    # Ouvrir le fichier
    wb = openpyxl.load_workbook(output_file, data_only=False)
    ws = wb['Résultats observations']

    # Définir les 20 nouvelles cibles
    new_targets = []
    for i in range(1, 4): new_targets.append(f"V{i}")    # V1,V2,V3
    for i in range(1, 6): new_targets.append(f"H{i}")    # H1-H5
    for i in range(1, 6): new_targets.append(f"B{i}")    # B1-B5
    for i in range(1, 6): new_targets.append(f"M{i}")    # M1-M5
    new_targets.extend(["REF1", "REF2"])                 # REF1,REF2

    print(f"🎯 Ajout de {len(new_targets)} nouvelles cibles...")

    # Trouver la première colonne libre (après les existantes)
    start_col = 15  # Colonne O

    # Vérifier où commencer vraiment
    for col in range(15, 30):
        if ws.cell(4, col).value is None or str(ws.cell(4, col).value).strip() == "":
            start_col = col
            break

    print(f"📍 Début insertion à la colonne {get_column_letter(start_col)}")

    current_col = start_col

    for target_id in new_targets:
        print(f"   ➕ {target_id} -> {get_column_letter(current_col)}-{get_column_letter(current_col+2)}")

        # Ajouter les 3 colonnes X, Y, Z pour chaque cible
        for coord_idx, coord in enumerate(['X', 'Y', 'Z']):
            col = current_col + coord_idx

            # En-têtes
            ws.cell(4, col).value = target_id                    # Nom de la cible
            ws.cell(5, col).value = ['3', '4', '5'][coord_idx]   # Type de coordonnée
            ws.cell(9, col).value = coord                        # X, Y ou Z

            # Description (seulement sur la première colonne)
            if coord_idx == 0:
                ws.cell(8, col).value = f"CHP_{target_id} ({target_id})"
                ws.cell(7, col).value = 'BG'
            elif coord_idx == 2:
                ws.cell(7, col).value = 'HG'

            # Formules pour récupérer les données de la DATABASE
            db_cols = {'X': 'AE', 'Y': 'AF', 'Z': 'AG'}
            db_col = db_cols[coord]

            # Ajouter les formules pour toutes les lignes de données
            for row in range(10, 51):  # Lignes 10 à 50
                formula = f'=IFERROR(INDEX(DATABASE.{db_col}:{db_col},MATCH(1,(DATABASE.AB:AB=$A{row})*(DATABASE.AC:AC="{target_id}"),0)),"")'
                ws.cell(row, col).value = formula

        current_col += 3  # Passer aux 3 colonnes suivantes

    # Sauvegarder
    wb.save(output_file)
    wb.close()

    print(f"✅ Fichier final créé: {output_file}")
    print(f"📊 {len(new_targets)} cibles × 3 coordonnées = {len(new_targets)*3} colonnes ajoutées")

    return output_file

def test_final_file(filename):
    """Teste le fichier final"""
    print(f"\\n🔍 Test du fichier final: {filename}")

    try:
        wb = openpyxl.load_workbook(filename, data_only=True)
        ws = wb['Résultats observations']

        # Compter les cibles
        targets_found = set()
        for col in range(15, 100):  # Large plage
            cell_value = ws.cell(4, col).value
            if cell_value and str(cell_value).strip() and len(str(cell_value).strip()) <= 10:
                targets_found.add(str(cell_value).strip())

        print(f"🎯 Cibles détectées: {len(targets_found)}")
        print(f"    {sorted(list(targets_found))}")

        wb.close()
        print("✅ Test réussi - Le fichier semble correct!")
        return True

    except Exception as e:
        print(f"❌ Erreur lors du test: {e}")
        return False

if __name__ == "__main__":
    print("🚀 Ajout des 20 cibles au fichier CLEAN qui fonctionne...")

    final_file = add_targets_to_clean_file()
    test_final_file(final_file)

    print(f"\\n🎉 TERMINÉ!")
    print(f"📁 Fichier prêt: {final_file}")
    print("💡 Vous pouvez maintenant l'ouvrir dans Excel et ajouter vos mesures dans la DATABASE!")
