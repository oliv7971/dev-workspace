#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SCRIPT DE TEST - Simulation du système de gestion des mesures
==============================================================

Ce script crée un fichier Excel de test pour vérifier le fonctionnement
des scripts de reset et d'ajout sans toucher au fichier réel.

Il génère :
- Un onglet DATABASE avec mesures multiples par jour
- Des onglets vides à remplir
- Des données de test réalistes
"""

import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
import random

FICHIER_TEST = "TEST_AUSCULTATION.xlsm"

def creer_fichier_test():
    """Crée un fichier Excel de test avec DATABASE"""
    
    print("🧪 CRÉATION DU FICHIER DE TEST")
    print("="*60)
    
    # Créer le classeur
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Supprimer la feuille par défaut
    
    # === ONGLET DATABASE ===
    ws_db = wb.create_sheet("DATABASE")
    
    # En-têtes ligne 12 (comme dans le fichier réel)
    ws_db.cell(12, 28).value = "DATE"    # AB
    ws_db.cell(12, 29).value = "NO"      # AC
    ws_db.cell(12, 30).value = "NO"      # AD
    ws_db.cell(12, 31).value = "X"       # AE
    ws_db.cell(12, 32).value = "Y"       # AF
    ws_db.cell(12, 33).value = "Z"       # AG
    
    print("📋 Création de l'onglet DATABASE...")
    
    # Générer des données de test
    cibles = ['V1', 'V2', 'V3', 'H1', 'H2', 'REF1', 'REF2']
    date_debut = datetime(2024, 8, 28)
    
    row = 14  # Première ligne de données
    
    # Campagnes normales (1 mesure par jour)
    print("   📅 Campagnes normales (1 mesure/jour)...")
    for jour in range(0, 30):  # 30 jours
        date_mesure = date_debut + timedelta(days=jour)
        
        for cible in cibles:
            x = 823147.64 + random.uniform(-0.01, 0.01)
            y = 1091710.08 + random.uniform(-0.01, 0.01)
            z = -121.56 + random.uniform(-0.01, 0.01)
            
            ws_db.cell(row, 28).value = date_mesure.strftime("%Y-%m-%d %H:%M:%S")
            ws_db.cell(row, 29).value = cible
            ws_db.cell(row, 30).value = f"{cible}.Test"
            ws_db.cell(row, 31).value = round(x, 4)
            ws_db.cell(row, 32).value = round(y, 4)
            ws_db.cell(row, 33).value = round(z, 4)
            row += 1
    
    nb_mesures_normales = row - 14
    
    # Campagnes ACCÉLÉRÉES (3 mesures par jour) - SIMULATION CRISE
    print("   🚨 Simulation surveillance accélérée (3 mesures/jour)...")
    date_crise = date_debut + timedelta(days=30)
    
    for jour in range(0, 3):  # 3 jours de crise
        # Mesure du matin (08:00)
        date_mesure_matin = date_crise + timedelta(days=jour, hours=8)
        
        # Mesure après-midi (14:00)
        date_mesure_midi = date_crise + timedelta(days=jour, hours=14)
        
        # Mesure soir (20:00)
        date_mesure_soir = date_crise + timedelta(days=jour, hours=20)
        
        for date_mesure in [date_mesure_matin, date_mesure_midi, date_mesure_soir]:
            for cible in cibles:
                x = 823147.64 + random.uniform(-0.01, 0.01)
                y = 1091710.08 + random.uniform(-0.01, 0.01)
                z = -121.56 + random.uniform(-0.01, 0.01)
                
                ws_db.cell(row, 28).value = date_mesure.strftime("%Y-%m-%d %H:%M:%S")
                ws_db.cell(row, 29).value = cible
                ws_db.cell(row, 30).value = f"{cible}.Test"
                ws_db.cell(row, 31).value = round(x, 4)
                ws_db.cell(row, 32).value = round(y, 4)
                ws_db.cell(row, 33).value = round(z, 4)
                row += 1
    
    nb_mesures_crise = row - 14 - nb_mesures_normales
    nb_total = row - 14
    
    print(f"   ✅ {nb_mesures_normales} mesures normales générées")
    print(f"   ✅ {nb_mesures_crise} mesures accélérées générées")
    print(f"   ✅ TOTAL : {nb_total} mesures")
    
    # === ONGLETS À REMPLIR ===
    print("\n📋 Création des onglets à remplir...")
    
    onglets = [
        "Résultats observations",
        "OBSERVATIONS",
        "PROJECTION",
        "EVOLUTIONS XYZ",
        "ÉVOLUTION PM H V",
        "cordes ref1",
        "cordes ref2",
        "EVOLUTIONS CORDES REF1",
        "EVOLUTIONS CORDES REF2"
    ]
    
    for onglet in onglets:
        ws = wb.create_sheet(onglet)
        
        # En-têtes basiques
        ws.cell(8, 1).value = "Date"
        ws.cell(8, 2).value = "Jours"
        
        # Quelques colonnes de données (vides pour l'instant)
        for col in range(3, 10):
            ws.cell(8, col).value = f"Col{col}"
        
        print(f"   ✅ {onglet}")
    
    # === SAUVEGARDER ===
    print(f"\n💾 Sauvegarde du fichier test...")
    wb.save(FICHIER_TEST)
    print(f"✅ Fichier créé : {FICHIER_TEST}")
    
    print("\n" + "="*60)
    print("✅ FICHIER DE TEST PRÊT")
    print("="*60)
    print(f"\n📊 Contenu :")
    print(f"   - DATABASE : {nb_total} mesures ({len(cibles)} cibles)")
    print(f"   - Période normale : {nb_mesures_normales // len(cibles)} jours (1 mesure/jour)")
    print(f"   - Période critique : {nb_mesures_crise // (len(cibles) * 3)} jours (3 mesures/jour)")
    print(f"   - {len(onglets)} onglets vides à remplir")
    
    print(f"\n🧪 Tests recommandés :")
    print(f"   1. Modifier reset_et_mise_a_jour_complet.py")
    print(f"      FICHIER_EXCEL = '{FICHIER_TEST}'")
    print(f"   2. Exécuter : python reset_et_mise_a_jour_complet.py")
    print(f"   3. Vérifier dans Excel les dates avec heures")
    print(f"   4. Tester : python ajouter_nouvelles_mesures.py")


def afficher_infos_database():
    """Affiche les infos du fichier DATABASE créé"""
    
    print("\n📊 ANALYSE DU FICHIER DE TEST")
    print("="*60)
    
    try:
        wb = openpyxl.load_workbook(FICHIER_TEST)
        ws_db = wb["DATABASE"]
        
        dates = []
        for row in range(14, ws_db.max_row + 1):
            date_val = ws_db.cell(row, 28).value
            if date_val:
                dates.append(date_val)
        
        print(f"✅ {len(dates)} mesures trouvées")
        
        # Grouper par date (sans heure)
        dates_uniquement = [d.split()[0] if isinstance(d, str) else d.strftime("%Y-%m-%d") for d in dates]
        from collections import Counter
        compteur = Counter(dates_uniquement)
        
        print(f"\n📅 Dates avec mesures multiples :")
        for date, count in compteur.most_common(10):
            if count > len(['V1', 'V2', 'V3', 'H1', 'H2', 'REF1', 'REF2']):  # Plus qu'une mesure/cible
                print(f"   {date} : {count // 7} mesures par cible")
        
        wb.close()
        
    except FileNotFoundError:
        print(f"❌ Fichier {FICHIER_TEST} non trouvé")
        print("   Exécutez d'abord : python test_systeme_mesures.py create")


if __name__ == "__main__":
    import sys
    
    print("\n" + "🧪 " * 30)
    print("SCRIPT DE TEST - Système de gestion des mesures")
    print("🧪 " * 30 + "\n")
    
    if len(sys.argv) > 1 and sys.argv[1] == "info":
        afficher_infos_database()
    else:
        creer_fichier_test()
        print("\n💡 Pour voir les détails : python test_systeme_mesures.py info")
