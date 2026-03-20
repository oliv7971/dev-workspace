#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
VISUALISEUR DE DATES DATABASE
==============================

Script d'analyse rapide pour visualiser les dates présentes dans DATABASE
et détecter les mesures multiples par jour.

Utile pour :
- Vérifier rapidement le contenu de DATABASE avant un reset
- Détecter les périodes de surveillance accélérée
- Identifier les dates manquantes ou problématiques
"""

import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime
from collections import defaultdict, Counter

FICHIER_EXCEL = "01_51-B-GER-HA21_3-PM_135-25-05-21.xlsm"

def analyser_database():
    """Analyse complète de l'onglet DATABASE"""
    
    print("\n📊 ANALYSE DATABASE")
    print("="*70)
    print(f"📂 Fichier : {FICHIER_EXCEL}\n")
    
    try:
        wb = openpyxl.load_workbook(FICHIER_EXCEL, data_only=True)
    except FileNotFoundError:
        print(f"❌ Fichier non trouvé : {FICHIER_EXCEL}")
        return
    except Exception as e:
        print(f"❌ Erreur : {e}")
        return
    
    if "DATABASE" not in wb.sheetnames:
        print("❌ Onglet DATABASE non trouvé")
        wb.close()
        return
    
    ws_db = wb["DATABASE"]
    
    # Trouver la colonne DATE
    date_col = None
    no_col = None
    
    for col in range(1, ws_db.max_column + 1):
        cell_value = ws_db.cell(12, col).value
        if cell_value:
            if str(cell_value).upper() == "DATE":
                date_col = col
            elif str(cell_value).upper() == "NO":
                if no_col is None:  # Première colonne NO
                    no_col = col
    
    if date_col is None:
        print("❌ Colonne DATE non trouvée")
        wb.close()
        return
    
    print(f"✅ Colonne DATE : {get_column_letter(date_col)}")
    if no_col:
        print(f"✅ Colonne NO (cibles) : {get_column_letter(no_col)}")
    
    # Extraire toutes les mesures
    print(f"\n📅 Extraction des mesures...")
    
    mesures = []
    for row in range(14, ws_db.max_row + 1):
        date_val = ws_db.cell(row, date_col).value
        cible_val = ws_db.cell(row, no_col).value if no_col else "?"
        
        if date_val:
            try:
                if isinstance(date_val, datetime):
                    dt = date_val
                elif isinstance(date_val, str):
                    for fmt in ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y"]:
                        try:
                            dt = datetime.strptime(date_val, fmt)
                            break
                        except ValueError:
                            continue
                    else:
                        continue
                
                mesures.append({
                    'datetime': dt,
                    'date': dt.date(),
                    'heure': dt.time(),
                    'cible': cible_val,
                    'ligne': row
                })
            except Exception:
                pass
    
    if not mesures:
        print("❌ Aucune mesure trouvée")
        wb.close()
        return
    
    print(f"✅ {len(mesures)} mesures trouvées\n")
    
    # Trier par date
    mesures.sort(key=lambda m: m['datetime'])
    
    # === STATISTIQUES GÉNÉRALES ===
    print("📊 STATISTIQUES GÉNÉRALES")
    print("-" * 70)
    
    premiere_mesure = mesures[0]
    derniere_mesure = mesures[-1]
    
    print(f"📅 Première mesure : {premiere_mesure['datetime'].strftime('%d/%m/%Y %H:%M')}")
    print(f"📅 Dernière mesure  : {derniere_mesure['datetime'].strftime('%d/%m/%Y %H:%M')}")
    
    duree = (derniere_mesure['datetime'] - premiere_mesure['datetime']).days
    print(f"⏱️  Durée totale     : {duree} jours")
    
    # Cibles uniques
    cibles_uniques = set(m['cible'] for m in mesures)
    print(f"🎯 Cibles trouvées  : {len(cibles_uniques)} ({', '.join(sorted(cibles_uniques))})")
    
    # === ANALYSE PAR JOUR ===
    print(f"\n📅 ANALYSE PAR JOUR")
    print("-" * 70)
    
    mesures_par_jour = defaultdict(list)
    for m in mesures:
        mesures_par_jour[m['date']].append(m)
    
    nb_jours_total = len(mesures_par_jour)
    print(f"📆 Nombre de jours mesurés : {nb_jours_total}")
    
    # Compter mesures par cible par jour
    mesures_par_cible_par_jour = {}
    for date, mesures_jour in mesures_par_jour.items():
        cibles_ce_jour = [m['cible'] for m in mesures_jour]
        compteur_cibles = Counter(cibles_ce_jour)
        mesures_par_cible_par_jour[date] = compteur_cibles
    
    # Détecter les mesures multiples
    jours_mesures_multiples = {}
    for date, compteur in mesures_par_cible_par_jour.items():
        max_mesures = max(compteur.values())
        if max_mesures > 1:
            jours_mesures_multiples[date] = max_mesures
    
    if jours_mesures_multiples:
        print(f"\n🔄 MESURES MULTIPLES PAR JOUR : {len(jours_mesures_multiples)} jours")
        print("-" * 70)
        
        for date in sorted(jours_mesures_multiples.keys()):
            nb_mesures_max = jours_mesures_multiples[date]
            mesures_ce_jour = mesures_par_jour[date]
            heures = sorted(set(m['heure'].strftime('%H:%M') for m in mesures_ce_jour))
            
            print(f"📅 {date.strftime('%d/%m/%Y')} : {nb_mesures_max} mesures/cible")
            print(f"   ⏰ Heures : {', '.join(heures)}")
            
            # Détail par cible
            compteur = mesures_par_cible_par_jour[date]
            cibles_multi = {c: n for c, n in compteur.items() if n > 1}
            if cibles_multi:
                print(f"   🎯 Cibles : {', '.join(f'{c}({n}x)' for c, n in cibles_multi.items())}")
            print()
        
        print(f"⚠️ ATTENTION : {len(jours_mesures_multiples)} jours avec mesures multiples")
        print(f"   → Surveillance accélérée détectée")
    else:
        print(f"✅ Aucune mesure multiple par jour - Une seule mesure/cible/jour")
    
    # === CALENDRIER SIMPLIFIÉ ===
    print(f"\n📆 CALENDRIER DES MESURES (30 derniers jours)")
    print("-" * 70)
    
    jours_tries = sorted(mesures_par_jour.keys())
    derniers_jours = jours_tries[-30:] if len(jours_tries) > 30 else jours_tries
    
    for date in derniers_jours:
        nb_mesures_total = len(mesures_par_jour[date])
        nb_cibles = len(set(m['cible'] for m in mesures_par_jour[date]))
        compteur = mesures_par_cible_par_jour[date]
        max_mesures = max(compteur.values())
        
        symbole = "🔴" if max_mesures > 2 else "🟡" if max_mesures > 1 else "🟢"
        
        print(f"{symbole} {date.strftime('%d/%m/%Y')} : {nb_mesures_total:3d} mesures "
              f"({nb_cibles} cibles × {max_mesures} mesure{'s' if max_mesures > 1 else ''})")
    
    print()
    print("Légende : 🟢 = 1 mesure/jour  🟡 = 2 mesures/jour  🔴 = 3+ mesures/jour")
    
    # === RECOMMANDATIONS ===
    print(f"\n💡 RECOMMANDATIONS")
    print("-" * 70)
    
    if jours_mesures_multiples:
        print("⚠️ Mesures multiples détectées")
        print("   → Utilisez 'reset_et_mise_a_jour_complet.py' pour gérer correctement")
        print("   → Les dates incluront l'heure (ex: 26/11/2025 14:00)")
    else:
        print("✅ Une seule mesure par jour")
        print("   → 'ajouter_nouvelles_mesures.py' suffira pour les mises à jour")
    
    print(f"\n📋 {len(mesures)} mesures prêtes à être traitées")
    
    wb.close()
    
    print("\n" + "="*70)
    print("✅ ANALYSE TERMINÉE")
    print("="*70)


def comparer_avec_onglets():
    """Compare DATABASE avec les dates dans les onglets"""
    
    print("\n🔍 COMPARAISON DATABASE ↔ ONGLETS")
    print("="*70)
    
    try:
        wb = openpyxl.load_workbook(FICHIER_EXCEL, data_only=True)
    except FileNotFoundError:
        print(f"❌ Fichier non trouvé : {FICHIER_EXCEL}")
        return
    
    # Extraire dates DATABASE
    if "DATABASE" not in wb.sheetnames:
        print("❌ Onglet DATABASE non trouvé")
        wb.close()
        return
    
    ws_db = wb["DATABASE"]
    
    # Trouver colonne DATE
    date_col = None
    for col in range(1, ws_db.max_column + 1):
        if ws_db.cell(12, col).value and str(ws_db.cell(12, col).value).upper() == "DATE":
            date_col = col
            break
    
    if not date_col:
        print("❌ Colonne DATE non trouvée")
        wb.close()
        return
    
    dates_db = set()
    for row in range(14, ws_db.max_row + 1):
        date_val = ws_db.cell(row, date_col).value
        if date_val:
            try:
                if isinstance(date_val, datetime):
                    dates_db.add(date_val)
                elif isinstance(date_val, str):
                    for fmt in ["%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y"]:
                        try:
                            dates_db.add(datetime.strptime(date_val, fmt))
                            break
                        except ValueError:
                            continue
            except:
                pass
    
    print(f"📊 DATABASE : {len(dates_db)} dates uniques")
    
    # Vérifier l'onglet Résultats observations
    if "Résultats observations" in wb.sheetnames:
        ws_obs = wb["Résultats observations"]
        
        dates_obs = set()
        for row in range(10, ws_obs.max_row + 1):
            date_val = ws_obs.cell(row, 1).value
            if date_val:
                try:
                    if isinstance(date_val, datetime):
                        dates_obs.add(date_val)
                    elif isinstance(date_val, str):
                        for fmt in ["%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"]:
                            try:
                                dates_obs.add(datetime.strptime(date_val, fmt))
                                break
                            except ValueError:
                                continue
                except:
                    pass
        
        print(f"📋 Résultats observations : {len(dates_obs)} dates")
        
        manquantes = dates_db - dates_obs
        en_trop = dates_obs - dates_db
        
        if manquantes:
            print(f"\n⚠️ {len(manquantes)} dates dans DATABASE mais PAS dans l'onglet :")
            for dt in sorted(manquantes)[:5]:
                print(f"   - {dt.strftime('%d/%m/%Y %H:%M')}")
            if len(manquantes) > 5:
                print(f"   ... et {len(manquantes) - 5} autres")
        
        if en_trop:
            print(f"\n⚠️ {len(en_trop)} dates dans l'onglet mais PAS dans DATABASE :")
            for dt in sorted(en_trop)[:5]:
                print(f"   - {dt.strftime('%d/%m/%Y %H:%M')}")
            if len(en_trop) > 5:
                print(f"   ... et {len(en_trop) - 5} autres")
        
        if not manquantes and not en_trop:
            print("\n✅ Parfaite synchronisation entre DATABASE et onglets")
        else:
            print(f"\n💡 Recommandation : Exécuter 'reset_et_mise_a_jour_complet.py'")
    
    wb.close()
    print()


if __name__ == "__main__":
    import sys
    
    print("\n" + "📊 " * 30)
    print("VISUALISEUR DE DATES DATABASE")
    print("📊 " * 30 + "\n")
    
    if len(sys.argv) > 1 and sys.argv[1] == "compare":
        comparer_avec_onglets()
    else:
        analyser_database()
        
        print("\n💡 Pour comparer avec les onglets : python visualiseur_dates.py compare")
