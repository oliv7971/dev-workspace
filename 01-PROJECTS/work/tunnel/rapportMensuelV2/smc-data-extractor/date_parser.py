"""
Script pour comparer les méthodes de calcul périodique
"""

import openpyxl
from pathlib import Path
from datetime import datetime, date
import pandas as pd

def parse_date(date_value):
    """Parse date from Excel cell"""
    if date_value is None:
        return None
    
    if isinstance(date_value, (datetime, date)):
        return date_value.date() if hasattr(date_value, 'date') else date_value
    
    # Si c'est une chaîne, essayer de la parser
    if isinstance(date_value, str):
        try:
            # Format Excel typique: "2024-01-15 00:00:00"
            parsed = datetime.strptime(date_value.split()[0], '%Y-%m-%d')
            return parsed.date()
        except:
            try:
                # Autres formats possibles
                parsed = datetime.strptime(date_value, '%Y-%m-%d')
                return parsed.date()
            except:
                return None
    
    return None

def clean_numeric_value(value):
    """Nettoie une valeur numérique"""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        try:
            return float(value.replace(',', '.'))
        except:
            return None
    return None

def debug_one_file():
    """Debug un seul fichier pour comprendre les différences"""
    
    # Prendre un fichier d'exemple
    excel_path = Path(r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250901-rapport mensuel aout\1-tableaux\GGS\03_104-GGS_SMC_C003_PM001_25_08-25.xlsm")
    
    if not excel_path.exists():
        print(f"Fichier non trouvé: {excel_path}")
        return
    
    print(f"🔍 Analyse: {excel_path.name}")
    
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb['Convergences']
    
    # Trouver les en-têtes
    headers = {}
    for col in range(1, 20):  # Colonnes A à S
        header = ws.cell(9, col).value
        if header and header not in ['Date', 'Temps (j)', 'jours écoulés', 'Distance au front']:
            headers[col] = str(header).strip()
    
    print(f"📋 Métriques trouvées: {list(headers.values())}")
    
    # Extraire toutes les données
    all_data = []
    print(f"🔍 Analysant les lignes 10 à {ws.max_row}...")
    
    for row in range(10, ws.max_row + 1):
        date_value = ws.cell(row, 2).value  # Colonne B = Date
        parsed_date = parse_date(date_value)
        
        if parsed_date is None:
            continue
            
        row_data = {'date': parsed_date, 'row': row}
        
        for col_idx, metric in headers.items():
            cell_value = ws.cell(row, col_idx).value
            clean_value = clean_numeric_value(cell_value)
            if clean_value is not None:
                row_data[metric] = clean_value
        
        if len(row_data) > 2:  # Date + row + au moins 1 métrique
            all_data.append(row_data)
    
    if not all_data:
        print("❌ Aucune donnée trouvée")
        return
    
    # Trier par date
    all_data.sort(key=lambda x: x['date'])
    
    print(f"📊 {len(all_data)} lignes de données trouvées")
    print(f"📅 Période: {all_data[0]['date']} à {all_data[-1]['date']}")
    
    # Regarder quelques lignes récentes pour débugger
    print("\n🔍 Dernières lignes de données:")
    for row_data in all_data[-5:]:
        print(f"  Ligne {row_data['row']}: Date={row_data['date']}, BG={row_data.get('BG', 'N/A')}")
    
    # Chercher un second tableau dans la feuille
    print(f"\n🔍 Recherche d'autres tableaux dans la feuille...")
    for row in range(70, min(ws.max_row + 1, 100)):
        for col in range(1, 15):
            value = ws.cell(row, col).value
            if value and isinstance(value, str) and ('convergence' in value.lower() or 'bg' in value.lower() or 'temps' in value.lower()):
                print(f"  Ligne {row}, Col {col}: '{value}'")
                
                # Regarder autour de cette cellule
                print(f"    Contexte ligne {row}:")
                for c in range(1, 15):
                    ctx_value = ws.cell(row, c).value
                    print(f"      Col {c}: '{ctx_value}'")
                
                # Regarder quelques lignes après
                print(f"    Données suivantes:")
                for r in range(row+1, min(row+5, ws.max_row+1)):
                    row_values = []
                    for c in range(1, 15):
                        val = ws.cell(r, c).value
                        row_values.append(str(val)[:10] if val else "")
                    print(f"      Ligne {r}: {row_values}")    # Définir la période d'août 2025
    month_start = date(2025, 8, 1)
    month_end = date(2025, 8, 31)
    
    # Filtrer les données
    dates_in_month = [row for row in all_data if month_start <= row['date'] <= month_end]
    dates_before_month = [row for row in all_data if row['date'] < month_start]
    
    print(f"📈 Données dans le mois: {len(dates_in_month)}")
    print(f"📈 Données avant le mois: {len(dates_before_month)}")
    
    # Analyser la métrique BG
    metric = 'BG'
    if metric in headers.values():
        print(f"\n🎯 Analyse détaillée pour {metric}:")
        
        # Méthode 1: Validateur (dernière ligne - avant-dernière ligne)
        last_row = all_data[-1]
        prev_row = all_data[-2] if len(all_data) > 1 else None
        
        cumulative_validator = last_row.get(metric)
        prev_value_validator = prev_row.get(metric) if prev_row else None
        periodic_validator = None
        if cumulative_validator is not None and prev_value_validator is not None:
            periodic_validator = cumulative_validator - prev_value_validator
        
        print(f"📊 Méthode VALIDATEUR:")
        print(f"   Dernière valeur (ligne {last_row['row']}): {cumulative_validator}")
        print(f"   Valeur précédente (ligne {prev_row['row'] if prev_row else 'N/A'}): {prev_value_validator}")
        print(f"   Périodique calculé: {periodic_validator}")
        
        # Méthode 2: Extracteur original (filtrage temporel)
        last_in_month_value = None
        last_before_month_value = None
        
        for row in reversed(dates_in_month):
            if metric in row and row[metric] is not None:
                last_in_month_value = row[metric]
                last_in_month_date = row['date']
                last_in_month_row = row['row']
                break
        
        for row in reversed(dates_before_month):
            if metric in row and row[metric] is not None:
                last_before_month_value = row[metric]
                last_before_month_date = row['date']
                last_before_month_row = row['row']
                break
        
        periodic_extracteur = None
        if last_in_month_value is not None and last_before_month_value is not None:
            periodic_extracteur = last_in_month_value - last_before_month_value
        
        cumulative_extracteur = None
        for row in reversed(all_data):
            if metric in row and row[metric] is not None:
                cumulative_extracteur = row[metric]
                break
        
        print(f"📊 Méthode EXTRACTEUR:")
        print(f"   Dernière valeur dans le mois (ligne {last_in_month_row if 'last_in_month_row' in locals() else 'N/A'}): {last_in_month_value}")
        print(f"   Dernière valeur avant le mois (ligne {last_before_month_row if 'last_before_month_row' in locals() else 'N/A'}): {last_before_month_value}")
        print(f"   Périodique calculé: {periodic_extracteur}")
        print(f"   Cumulative: {cumulative_extracteur}")
        
        # Charger le CSV pour comparaison
        csv_path = Path("C:/temp/smc_output/Convergences_SMC.csv")
        if csv_path.exists():
            df = pd.read_csv(csv_path, encoding='utf-8-sig')
            csv_row = df[(df['galerie'] == 'GGS') & (df['section'] == 'C003') & (df['metric'] == metric)]
            if not csv_row.empty:
                csv_periodic = csv_row.iloc[0]['periodic_mm']
                csv_cumulative = csv_row.iloc[0]['cumulative_mm']
                print(f"📊 VALEURS CSV:")
                print(f"   Périodique: {csv_periodic}")
                print(f"   Cumulative: {csv_cumulative}")
                
                print(f"\n🎯 COMPARAISON:")
                print(f"   Validateur vs CSV (périodique): {periodic_validator} vs {csv_periodic}")
                print(f"   Extracteur vs CSV (périodique): {periodic_extracteur} vs {csv_periodic}")
                print(f"   Cumulative (tous): {cumulative_validator} vs {cumulative_extracteur} vs {csv_cumulative}")

if __name__ == "__main__":
    debug_one_file()
