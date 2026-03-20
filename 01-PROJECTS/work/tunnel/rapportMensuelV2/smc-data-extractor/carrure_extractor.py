#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Extracteur spécialisé pour les fichiers Excel d'auscultation de carrure
Complément au système SMC existant
"""

import openpyxl
import pandas as pd
from pathlib import Path
from datetime import datetime, date
import logging
from typing import Dict, List, Optional, Tuple, Any

class CarrureExtractor:
    """Extracteur pour fichiers d'auscultation de carrure"""

    def __init__(self, log_callback=None):
        self.log_callback = log_callback or print
        self.extracted_data = {
            'carrure_displacements': [],
            'cintres_displacements': [],
            'point_positions': [],
            'thresholds': [],
            'precision_measures': []
        }

    def log(self, message: str):
        """Fonction de logging"""
        if self.log_callback:
            self.log_callback(message)

    def parse_date(self, date_value) -> Optional[date]:
        """Parse une date depuis une cellule Excel"""
        if date_value is None:
            return None

        if isinstance(date_value, (datetime, date)):
            return date_value.date() if hasattr(date_value, 'date') else date_value

        if isinstance(date_value, str):
            try:
                parsed = datetime.strptime(date_value.split()[0], '%Y-%m-%d')
                return parsed.date()
            except:
                try:
                    parsed = datetime.strptime(date_value, '%Y-%m-%d')
                    return parsed.date()
                except:
                    return None
        return None

    def clean_numeric_value(self, value) -> Optional[float]:
        """Nettoie une valeur numérique"""
        if value is None:
            return None
        if isinstance(value, (int, float)):
            return float(value) if not pd.isna(value) else None
        if isinstance(value, str):
            try:
                # Gérer les erreurs Excel comme #N/A
                if value.startswith('#'):
                    return None
                return float(value.replace(',', '.'))
            except:
                return None
        return None

    def extract_carrure_displacements(self, worksheet) -> List[Dict]:
        """Extrait les données de l'onglet 'Déplacements carrure'"""

        self.log("🔄 Extraction des déplacements de carrure...")
        data = []

        try:
            # Trouver la ligne d'en-têtes (ligne 10 d'après l'analyse)
            header_row = 10

            # Identifier les colonnes importantes
            date_col = 2  # Date
            jours_col = 1  # Jours écoulés
            front_col = 3  # Front

            # Colonnes de déplacements pour chaque point (groupes de colonnes)
            # Chaque point a généralement DPM, DH, DZ
            point_groups = []
            for col in range(4, worksheet.max_column + 1, 3):  # Groupes de 3 colonnes
                if col + 2 <= worksheet.max_column:
                    point_groups.append((col, col + 1, col + 2))  # DPM, DH, DZ

            # Extraire les données ligne par ligne à partir de la ligne 12
            for row in range(12, worksheet.max_row + 1):
                date_value = worksheet.cell(row, date_col).value
                parsed_date = self.parse_date(date_value)

                if parsed_date is None:
                    continue

                jours_value = self.clean_numeric_value(worksheet.cell(row, jours_col).value)
                front_value = self.clean_numeric_value(worksheet.cell(row, front_col).value)

                # Données de base
                row_data = {
                    'date': parsed_date,
                    'jours_ecoules': jours_value,
                    'front': front_value,
                    'file_type': 'carrure',
                    'sheet_name': 'Déplacements carrure'
                }

                # Ajouter les déplacements par point
                for i, (dpm_col, dh_col, dz_col) in enumerate(point_groups):
                    point_name = f"point_{i+1}"  # Nom générique, sera affiné

                    dpm_val = self.clean_numeric_value(worksheet.cell(row, dpm_col).value)
                    dh_val = self.clean_numeric_value(worksheet.cell(row, dh_col).value)
                    dz_val = self.clean_numeric_value(worksheet.cell(row, dz_col).value)

                    row_data[f'{point_name}_dpm'] = dpm_val
                    row_data[f'{point_name}_dh'] = dh_val
                    row_data[f'{point_name}_dz'] = dz_val

                data.append(row_data)

        except Exception as e:
            self.log(f"❌ Erreur extraction déplacements carrure: {e}")

        self.log(f"✅ {len(data)} lignes de déplacements carrure extraites")
        return data

    def extract_cintres_displacements(self, worksheet) -> List[Dict]:
        """Extrait les données de l'onglet 'Déplacements cintres-longerons'"""

        self.log("🔄 Extraction des déplacements cintres-longerons...")
        data = []

        try:
            # Structure similaire aux déplacements carrure
            header_row = 10
            date_col = 2
            jours_col = 1
            front_col = 3

            # Extraire les données
            for row in range(12, worksheet.max_row + 1):
                date_value = worksheet.cell(row, date_col).value
                parsed_date = self.parse_date(date_value)

                if parsed_date is None:
                    continue

                jours_value = self.clean_numeric_value(worksheet.cell(row, jours_col).value)
                front_value = self.clean_numeric_value(worksheet.cell(row, front_col).value)

                row_data = {
                    'date': parsed_date,
                    'jours_ecoules': jours_value,
                    'front': front_value,
                    'file_type': 'carrure',
                    'sheet_name': 'Déplacements cintres-longerons'
                }

                # Ajouter les mesures de cintres (colonnes à partir de 4)
                for col in range(4, min(20, worksheet.max_column + 1)):  # Limiter pour éviter trop de colonnes
                    val = self.clean_numeric_value(worksheet.cell(row, col).value)
                    if val is not None:
                        row_data[f'cintre_col_{col}'] = val

                data.append(row_data)

        except Exception as e:
            self.log(f"❌ Erreur extraction cintres: {e}")

        self.log(f"✅ {len(data)} lignes de déplacements cintres extraites")
        return data

    def extract_from_carrure_file(self, file_path: Path, target_month: str = None) -> Dict:
        """Extrait toutes les données d'un fichier de carrure"""

        self.log(f"🔍 Analyse fichier carrure: {file_path.name}")

        if not file_path.exists():
            self.log(f"❌ Fichier non trouvé: {file_path}")
            return {}

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)

            # Extraire les déplacements carrure
            if 'Déplacements carrure' in wb.sheetnames:
                carrure_data = self.extract_carrure_displacements(wb['Déplacements carrure'])
                self.extracted_data['carrure_displacements'].extend(carrure_data)

            # Extraire les déplacements cintres
            if 'Déplacements cintres-longerons' in wb.sheetnames:
                cintres_data = self.extract_cintres_displacements(wb['Déplacements cintres-longerons'])
                self.extracted_data['cintres_displacements'].extend(cintres_data)

            # TODO: Ajouter extraction des autres onglets si nécessaire
            # - Position des points auscultés
            # - seuils
            # - Précision des mesures

            wb.close()

            self.log(f"✅ Extraction fichier carrure terminée")
            return self.extracted_data

        except Exception as e:
            self.log(f"❌ Erreur lors de l'extraction: {e}")
            return {}

    def save_to_csv(self, output_folder: Path):
        """Sauvegarde les données extraites en CSV"""

        output_folder.mkdir(exist_ok=True)

        # Sauvegarder déplacements carrure
        if self.extracted_data['carrure_displacements']:
            df_carrure = pd.DataFrame(self.extracted_data['carrure_displacements'])
            carrure_file = output_folder / "carrure_deplacements.csv"
            df_carrure.to_csv(carrure_file, index=False, encoding='utf-8-sig')
            self.log(f"💾 Déplacements carrure sauvés: {carrure_file}")

        # Sauvegarder déplacements cintres
        if self.extracted_data['cintres_displacements']:
            df_cintres = pd.DataFrame(self.extracted_data['cintres_displacements'])
            cintres_file = output_folder / "cintres_deplacements.csv"
            df_cintres.to_csv(cintres_file, index=False, encoding='utf-8-sig')
            self.log(f"💾 Déplacements cintres sauvés: {cintres_file}")

    def generate_summary_report(self, output_folder: Path):
        """Génère un rapport de synthèse"""

        report_file = output_folder / "carrure_summary_report.txt"

        with open(report_file, 'w', encoding='utf-8') as f:
            f.write("# RAPPORT D'EXTRACTION - AUSCULTATION DE CARRURE\n")
            f.write("=" * 60 + "\n\n")
            f.write(f"Date génération: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n")

            f.write("## DONNÉES EXTRAITES\n")
            f.write(f"- Déplacements carrure: {len(self.extracted_data['carrure_displacements'])} lignes\n")
            f.write(f"- Déplacements cintres: {len(self.extracted_data['cintres_displacements'])} lignes\n\n")

            # Période couverte
            if self.extracted_data['carrure_displacements']:
                dates = [row['date'] for row in self.extracted_data['carrure_displacements'] if row['date']]
                if dates:
                    f.write(f"## PÉRIODE COUVERTE\n")
                    f.write(f"- Du: {min(dates)}\n")
                    f.write(f"- Au: {max(dates)}\n")
                    f.write(f"- Nombre de jours: {len(set(dates))}\n\n")

        self.log(f"📋 Rapport de synthèse généré: {report_file}")

def run_carrure_extraction(root_folder: str, month: str, output_folder: str, options: Dict, log_callback=None) -> bool:
    """Fonction principale d'extraction pour fichiers de carrure"""

    extractor = CarrureExtractor(log_callback)

    try:
        root_path = Path(root_folder)
        output_path = Path(output_folder)

        # Chercher les fichiers de carrure (pattern spécifique)
        carrure_files = list(root_path.rglob("*carrure*.xlsx")) + list(root_path.rglob("*carrure*.xlsm"))

        if not carrure_files:
            extractor.log("⚠️  Aucun fichier de carrure trouvé")
            return False

        extractor.log(f"🎯 {len(carrure_files)} fichier(s) de carrure trouvé(s)")

        # Extraire chaque fichier
        for file_path in carrure_files:
            extractor.extract_from_carrure_file(file_path, month)

        # Sauvegarder les résultats
        if options.get('generate_csv', True):
            extractor.save_to_csv(output_path)

        # Générer le rapport
        extractor.generate_summary_report(output_path)

        return True

    except Exception as e:
        extractor.log(f"❌ Erreur extraction carrure: {e}")
        return False

# Test rapide
if __name__ == "__main__":
    test_file = Path(r"C:\data\11-CHANTIERS\BURE\10-ACTIVITES\Rapports d'activité\2025\250902-rapport mensuel septembre\1-tableaux\GGS\03_100-GGS-GRD-double carrure-XYZ-tableau Jour-25-09-10.xlsm")

    extractor = CarrureExtractor()
    result = extractor.extract_from_carrure_file(test_file)

    print(f"\n📊 RÉSULTAT TEST:")
    for key, data in result.items():
        print(f"   {key}: {len(data)} éléments")
