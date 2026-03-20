#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script autonome (sans ligne de commande) — ANNEXE DÉPLACEMENTS
--------------------------------------------------------------
- Ouvre un sélecteur de fichier pour choisir `deplacements_YYYY-MM.csv`
- Génère un fichier XLSX avec un onglet "Annexe Deplacements"
- Mise en forme par FICHIER, comme pour l'annexe convergences :

    NOM DU FICHIER
    (ligne vide)
    DEPLACEMENTS PERIODIQUES — DH
    [BG IG HG HD ID BD LH LI LB]
    [valeurs périodiques DH]
    (ligne vide)
    DEPLACEMENTS PERIODIQUES — DZ
    [BG IG HG HD ID BD LH LI LB]
    [valeurs périodiques DZ]
    (ligne vide)
    DEPLACEMENTS PERIODIQUES — DPM   (affiché seulement si présent)
    ...
    (ligne vide)
    DEPLACEMENTS CUMULES — DH
    ... idem …

- Sauvegarde par défaut : <dossier_du_CSV>\_recaps\annexe_deplacements.xlsx

Usage :
- Double-cliquez le fichier (si l'association .py est configurée) ou
- Clic droit → Ouvrir avec → Python
"""
from __future__ import annotations

import os
from pathlib import Path
import pandas as pd
from tkinter import Tk, filedialog, messagebox
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

POSITIONS = ["BG", "IG", "HG", "HD", "ID", "BD", "LH", "LI", "LB"]
KINDS = ["DH", "DZ", "DPM"]  # ordre d'affichage souhaité

# --- Helpers Excel ---

def _last_notna(series: pd.Series):
    s = series.dropna()
    return None if s.empty else s.iloc[-1]

def _values_for_kind(df_file: pd.DataFrame, kind: str, value_col: str):
    """Retourne une liste de 9 valeurs (ou "") pour le type donné (DH/DZ/DPM) dans l'ordre POSITIONS."""
    vals = []
    dsub = df_file[df_file["kind"].astype(str).str.upper().str.strip() == kind]
    for pos in POSITIONS:
        sub = dsub[dsub["position"].astype(str).str.upper().str.strip() == pos]
        if sub.empty:
            vals.append("")
            continue
        v = _last_notna(sub[value_col])
        vals.append("" if pd.isna(v) else float(v))
    return vals

def _autosize(ws) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for cell in ws[letter]:
            txt = "" if cell.value is None else str(cell.value)
            if len(txt) > max_len:
                max_len = len(txt)
        ws.column_dimensions[letter].width = min(max_len + 2, 40)


def _write_table(ws, start_row: int, title: str, headers: list[str], values: list):
    bold = Font(bold=True)
    big  = Font(bold=True, size=12)
    center = Alignment(horizontal="center")
    r = start_row

    # titre
    c = ws.cell(row=r, column=1, value=title); c.font = big
    r += 1

    # entêtes
    for j, h in enumerate(headers, start=1):
        cc = ws.cell(row=r, column=j, value=h); cc.font = bold; cc.alignment = center

    # valeurs
    for j, v in enumerate(values, start=1):
        cc = ws.cell(row=r + 1, column=j, value=v)
        if isinstance(v, (int, float)): cc.number_format = "0.00"
        cc.alignment = center

    return r + 3  # saute une ligne après le tableau


def build_annexe_from_csv(disp_csv: Path, out_xlsx: Path) -> Path:
    df = pd.read_csv(disp_csv, encoding="utf-8-sig")
    needed = {"source_file", "kind", "position", "periodic_mm", "cumulative_mm"}
    missing = needed - set(df.columns)
    if missing:
        raise ValueError(f"Colonnes manquantes dans {disp_csv}: {sorted(missing)}")

    wb = Workbook()
    ws = wb.active
    ws.title = "Annexe Deplacements"

    row = 1
    title_font = Font(bold=True, size=14)

    # Grouper par fichier (toutes feuilles confondues)
    for src, grp in df.groupby("source_file", sort=False):
        filename = os.path.basename(src)
        ws.cell(row=row, column=1, value=filename).font = title_font
        row += 2

        # PERIODIQUES
        for kind in KINDS:
            if not (grp["kind"].astype(str).str.upper() == kind).any():
                continue  # ne pas afficher les sections vides
            vals = _values_for_kind(grp, kind, "periodic_mm")
            row = _write_table(ws, row, f"DEPLACEMENTS PERIODIQUES — {kind}", POSITIONS, vals)

        # CUMULES
        for kind in KINDS:
            if not (grp["kind"].astype(str).str.upper() == kind).any():
                continue
            vals = _values_for_kind(grp, kind, "cumulative_mm")
            row = _write_table(ws, row, f"DEPLACEMENTS CUMULES — {kind}", POSITIONS, vals)

        row += 2  # séparation avant le fichier suivant

    _autosize(ws)
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx

# --- UI minimale ---

def main():
    from tkinter import Tk
    root = Tk(); root.withdraw()

    messagebox.showinfo("Annexe déplacements", "Choisissez le fichier deplacements_YYYY-MM.csv")
    path = filedialog.askopenfilename(
        title="Choisir deplacements_YYYY-MM.csv",
        filetypes=[("CSV", "*.csv")]
    )
    if not path:
        return

    disp_csv = Path(path)
    out_xlsx = disp_csv.parent / "_recaps" / "annexe_deplacements.xlsx"

    try:
        result = build_annexe_from_csv(disp_csv, out_xlsx)
    except Exception as e:
        messagebox.showerror("Erreur", f"Impossible de créer l'annexe:\n{e}")
        return

    messagebox.showinfo("OK", f"Annexe créée :\n{result}")
    try:
        os.startfile(str(result))
    except Exception:
        pass

if __name__ == "__main__":
    main()
