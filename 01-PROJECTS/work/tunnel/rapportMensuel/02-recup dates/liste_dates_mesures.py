import openpyxl
from openpyxl.utils import column_index_from_string

def formater_dates(fichier_excel, feuille_nom):
    try:
        # Charger le fichier Excel
        workbook = openpyxl.load_workbook(fichier_excel)
        sheet = workbook[feuille_nom]

        # Demander à l'utilisateur de spécifier la plage de cellules
        plage = input("Veuillez entrer la plage de cellules (par exemple, F2:N18) : ")

        # Extraire les informations de la plage
        start_col, start_row, end_col, end_row = parse_range(plage)

        for row in sheet.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
            dates = []
            for cell in row:
                if cell.value is not None:
                    # Assurez-vous que la valeur est un objet date ou peut être convertie en date
                    try:
                        date = cell.value if isinstance(cell.value, str) else cell.value.strftime('%d/%m/%Y')
                        dates.append(date)
                    except AttributeError:
                        # Si la conversion échoue, ajoutez simplement la valeur telle quelle
                        dates.append(str(cell.value))

            output = formater_ligne_dates(dates)

            # Écrire le résultat dans la cellule à droite
            sheet.cell(row=row[0].row, column=end_col + 1, value=output)

        # Sauvegarder le fichier Excel
        workbook.save(fichier_excel)
        print("Les dates ont été formatées et enregistrées avec succès.")

    except Exception as e:
        print(f"Une erreur s'est produite : {e}")

def parse_range(plage):
    try:
        start, end = plage.split(':')
        start_col = column_index_from_string(''.join(filter(str.isalpha, start)))
        start_row = int(''.join(filter(str.isdigit, start)))
        end_col = column_index_from_string(''.join(filter(str.isalpha, end)))
        end_row = int(''.join(filter(str.isdigit, end)))
        return start_col, start_row, end_col, end_row
    except ValueError as e:
        raise ValueError(f"Format de plage invalide : {e}")

def formater_ligne_dates(dates):
    formatted_dates = []
    for i, date in enumerate(dates):
        if i == len(dates) - 1:
            formatted_dates.append(f"et le {date}")
        else:
            formatted_dates.append(f"le {date}")
    return ', '.join(formatted_dates)



# Exemple d'utilisation
fichier_excel = 'C:\\Temp\\1-tableaux\\_recaps\\_master_recap_2025-07.xlsx'
feuille_nom = 'Dates'
formater_dates(fichier_excel, feuille_nom)
