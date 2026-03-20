#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Diagnostic des Graphiques - Problème de Données Uniques
======================================================

Ce script analyse pourquoi les graphiques n'ont qu'une seule valeur par série
et propose des corrections pour afficher l'évolution temporelle.

Auteur: Système d'automatisation Excel
Date: Octobre 2025
"""

import openpyxl
import sys

def analyser_donnees_graphiques(fichier_excel):
    """Analyse les données utilisées dans les graphiques"""
    print("🔍 DIAGNOSTIC DES DONNÉES GRAPHIQUES")
    print("=" * 50)

    try:
        wb = openpyxl.load_workbook(fichier_excel)

        # Analyser l'onglet EVOLUTIONS XYZ pour voir les vraies données
        if 'EVOLUTIONS XYZ' in wb.sheetnames:
            ws = wb['EVOLUTIONS XYZ']
            print("\n📊 ANALYSE ONGLET: EVOLUTIONS XYZ")

            # Vérifier les lignes avec données
            print(f"  📏 Dimensions: {ws.max_row} lignes × {ws.max_column} colonnes")

            # Analyser les 15 premières lignes pour voir la structure
            print("\n📋 STRUCTURE DES DONNÉES:")
            for row in range(1, min(16, ws.max_row + 1)):
                row_data = []
                for col in range(1, min(11, ws.max_column + 1)):  # 10 premières colonnes
                    cell = ws.cell(row=row, column=col)
                    value = cell.value
                    if value is not None:
                        if isinstance(value, str) and len(value) > 15:
                            value = value[:15] + "..."
                        row_data.append(str(value))
                    else:
                        row_data.append("NULL")

                if any(val != "NULL" for val in row_data):
                    print(f"  Ligne {row:2d}: {' | '.join(row_data)}")

            # Compter les lignes avec vraies données (non nulles)
            lignes_avec_donnees = 0
            for row in range(10, ws.max_row + 1):  # À partir de ligne 10
                if ws.cell(row=row, column=1).value is not None:
                    lignes_avec_donnees += 1

            print(f"\n📈 LIGNES AVEC DONNÉES: {lignes_avec_donnees}")

            if lignes_avec_donnees <= 1:
                print("❌ PROBLÈME DÉTECTÉ: Seulement 1 ligne de données!")
                print("   → Les graphiques n'auront qu'un point chacun")
                print("   → Il faut des données temporelles multiples")

        # Analyser l'onglet Graphiques
        if 'Graphiques' in wb.sheetnames:
            ws_graph = wb['Graphiques']
            print(f"\n📊 ONGLET GRAPHIQUES: {len(ws_graph._charts)} graphique(s) détecté(s)")

            # Analyser chaque graphique
            for i, chart in enumerate(ws_graph._charts):
                print(f"\n📈 Graphique {i+1}:")
                print(f"   Titre: {chart.title}")
                print(f"   Type: {type(chart).__name__}")
                print(f"   Séries: {len(chart.series)}")

                for j, serie in enumerate(chart.series):
                    if hasattr(serie, 'values') and serie.values:
                        print(f"   Série {j+1}: {serie.values}")
                    if hasattr(serie, 'xValues') and serie.xValues:
                        print(f"   Données X: {serie.xValues}")

        wb.close()

    except Exception as e:
        print(f"❌ Erreur: {e}")
        import traceback
        traceback.print_exc()

def proposer_solutions():
    """Propose des solutions pour corriger le problème"""
    print("\n💡 SOLUTIONS PROPOSÉES")
    print("=" * 30)
    print("1. 📅 AJOUTER DONNÉES TEMPORELLES:")
    print("   → Créer plusieurs lignes de mesures avec dates")
    print("   → Une ligne = une campagne de mesure")
    print("   → Colonne A = Date de mesure")
    print()
    print("2. 🔄 CORRIGER LES PLAGES DE GRAPHIQUES:")
    print("   → Étendre les plages de données dans les graphiques")
    print("   → Inclure toutes les lignes de mesures historiques")
    print()
    print("3. 📊 GÉNÉRER DONNÉES D'EXEMPLE:")
    print("   → Créer des données temporelles de test")
    print("   → Simuler l'évolution des déplacements")

def main():
    if len(sys.argv) > 1:
        fichier = sys.argv[1]
    else:
        fichier = input("📂 Fichier Excel à analyser: ").strip()

    if not fichier:
        fichier = "AUSCULTATION CARREFOUR_TEMPLATE_v4_GRAPHIQUES_20_CIBLES_OPTIMISE.xlsx"

    analyser_donnees_graphiques(fichier)
    proposer_solutions()

if __name__ == "__main__":
    main()
